# -*- coding: utf-8 -*-
"""Import of files shaped like the real fastmag exports (made-up people and articles)."""
import base64
import io
from datetime import date, datetime, timedelta

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged
from odoo.tools.misc import xlsxwriter

ARTICLE_HEADER = ['site', 'réf article', 'EAN bar', None, 'catégor', 'Désignation', 'bu', 'marque', 'ligne', 'axe',
                  'sous axe', 'nature', None, 'Cont', 'Genre', None, 'statut', None, 'Prix', None, 'UN', None,
                  'Date creati', 'Fourn CODE', None, 'Fourn DES']
NEW_HEADER = ['CardNumbre', 'SurName', 'FirstName', 'ID', 'Adress', 'Street2', 'Street3', 'City', 'PostalCode',
              'Phone', 'Birthday', 'EntranceDate', 'CodeShop', 'CodeSeller', 'FirstPurchase', 'Mobile', 'Email',
              'Montant_TTC']
OLD_HEADER = ['compteur', 'numCliente', 'numClienteOLD', 'naissance', 'nom_x', 'prenom', 'patronyme', 'rue1', 'cp',
              'cp2', 'ville', 'villeAttache', 'CdB', 'NPAI', 'numVendeuse', 'dateExcel', 'dateExcel2', 'dateSup',
              'Seg1', 'Seg2', 'Seg3', 'Seg4', 'Segment', 'dateLastSegment', 'Tel2', 'CdB_H', 'CdB_date', 'SegLettre',
              'Rue2', 'Rue3', 'Champ1', 'sexe']
CA_HEADER = ['CodeMag', 'date', 'Client', 'Cartefidelite', 'Nom', 'Prenom', 'TOTALCA']
CODES_HEADER = ['CodeClient', 'Cartefidelite', 'Nom', 'datecreation', 'FidelitePoint', 'Catégorie']
DETAIL_HEADER = ['CodeMag', 'Vente', 'date', 'Utilisateur', 'Vendeur', 'Client', 'Cartefidelite', 'Nom', 'Prenom',
                 'Barcode', 'gencod', 'Reffournisseur', 'Designation', 'Axe', 'Lignes', 'Quantite', 'Prix', 'Remise',
                 'Total', 'Motif', 'Commentaire']


def xlsx(rows):
    out = io.BytesIO()
    workbook = xlsxwriter.Workbook(out, {'in_memory': True})
    sheet = workbook.add_worksheet('fastmag')
    date_format = workbook.add_format({'num_format': 'dd/mm/yyyy'})
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            if isinstance(value, datetime):
                sheet.write_datetime(r, c, value, date_format)
            elif value is not None:
                sheet.write(r, c, value)
    workbook.close()
    return base64.b64encode(out.getvalue())


def xlsx_sheets(sheets):
    """{sheet name: rows} → one workbook with several sheets."""
    out = io.BytesIO()
    workbook = xlsxwriter.Workbook(out, {'in_memory': True})
    date_format = workbook.add_format({'num_format': 'dd/mm/yyyy'})
    for name, rows in sheets.items():
        sheet = workbook.add_worksheet(name)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                if isinstance(value, datetime):
                    sheet.write_datetime(r, c, value, date_format)
                elif value is not None:
                    sheet.write(r, c, value)
    workbook.close()
    return base64.b64encode(out.getvalue())


def article(code, ean, name, axe, sub_axe, price, status='Actif ', supplier=('YVES', 'Y.ROCHER')):
    row = [None] * len(ARTICLE_HEADER)
    row[0], row[1], row[2], row[4], row[5], row[6], row[7], row[8] = \
        'DIST', code, ean, 'CAFV', name, 'YVES ROCHER', 'YVES ROCHER', 'CN3'
    row[9], row[10], row[11], row[16], row[18], row[20], row[22] = \
        axe, sub_axe, 'CORRECTEUR', status, price, 'UN', '16/09/2019'
    row[23], row[25] = supplier
    return row


@tagged('post_install', '-at_install')
class TestMigrationImport(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.store = cls.env['loyalty.store'].create({'name': 'YR Youg', 'code': 'YOUG', 'store_type': 'store',
                                                     'yr_fastmag_code': 'YR_YOUG'})

    def _import(self, rows, filename='fichier.xlsx', **options):
        wizard = self.env['yr.migration.import'].create(dict({'file': xlsx(rows), 'filename': filename}, **options))
        wizard.action_import()
        return wizard

    # ── articles ─────────────────────────────────────────────────────────
    def test_articles(self):
        empty = [None] * len(ARTICLE_HEADER)
        rows = [ARTICLE_HEADER, empty,
                article('T100102200', '9990005911601', 'Correcteur stick Zéro Défaut', 'MAQUILLAGE TEST', 'TEINT TEST', 17.349),
                empty, empty,
                article('T100102201', '9990005915821', 'Correcteur teinte 2', 'MAQUILLAGE TEST', 'TEINT TEST', 17.349),
                empty,
                article('T100200500', '9990005500500', 'Crème riche', 'SOIN TEST', 'HYDRA TEST', 38,
                        status='Inactif')]
        wizard = self._import(rows)
        self.assertEqual(wizard.file_type, 'article')
        Template = self.env['product.template']
        stick = Template.search([('default_code', '=', 'T100102200')])
        self.assertEqual((stick.barcode, stick.name, stick.list_price), ('9990005911601',
                                                                         'Correcteur stick Zéro Défaut', 17.349))
        self.assertTrue(stick.is_storable and stick.available_in_pos)
        # the axes are the categories: MAQUILLAGE › TEINT
        self.assertEqual(stick.pos_categ_ids.name, 'TEINT TEST')
        self.assertEqual(stick.pos_categ_ids.parent_id.name, 'MAQUILLAGE TEST')
        self.assertEqual((stick.yr_axe_id.name, stick.yr_sub_axe_id.name), ('MAQUILLAGE TEST', 'TEINT TEST'))
        self.assertTrue(stick.yr_axe_id.is_crm_family)
        self.assertEqual((stick.yr_category_code, stick.yr_line, stick.yr_status), ('CAFV', 'CN3', 'Actif'))
        self.assertEqual(stick.yr_create_date_legacy, date(2019, 9, 16))
        self.assertEqual(stick.seller_ids.partner_id.name, 'Y.ROCHER')
        self.assertEqual(stick.seller_ids.partner_id.ref, 'YVES')
        self.assertEqual(Template.search_count([('yr_sub_axe_id.name', '=', 'TEINT TEST')]), 2)
        self.assertEqual(self.env['pos.category'].search_count([('name', '=', 'TEINT TEST')]), 1)
        self.assertIn('Créés : <b>3</b>', wizard.result)
        # re-import: updated, nothing duplicated, archived option
        rows[2] = article('T100102200', '9990005911601', 'Correcteur stick', 'MAQUILLAGE TEST', 'TEINT TEST', 18.5)
        wizard = self._import(rows, archive_inactive=True)
        self.assertEqual(Template.search_count([('default_code', '=', 'T100102200')]), 1)
        self.assertEqual((stick.name, stick.list_price), ('Correcteur stick', 18.5))
        self.assertEqual(len(stick.seller_ids), 1)
        self.assertFalse(Template.with_context(active_test=False).search(
            [('default_code', '=', 'T100200500')]).active)
        self.assertIn('Mis à jour : <b>3</b>', wizard.result)

    def test_articles_imported_before_by_crm_screen(self):
        """Articles already created by Fidélité & CRM › Import articles: completed, not duplicated."""
        rows = [ARTICLE_HEADER,
                article('T100191452', '9990005648712', 'Brillance vinaigre', 'CAPILLAIRE TEST', 'SOIN CHEVEUX', 12),
                article('T100191954', '9990005648712', 'Capillaire brillance', 'CAPILLAIRE TEST', 'SOIN CHEVEUX', 12)]
        # articles created before by the CRM import (without the fastmag fields)
        Template = self.env['product.template']
        Template.create([{'name': 'Brillance vinaigre', 'default_code': 'T100191452', 'barcode': '9990005648712',
                          'is_storable': True},
                         {'name': 'Capillaire brillance', 'default_code': 'T100191954', 'is_storable': True}])
        first = Template.search([('default_code', '=', 'T100191452')])
        second = Template.search([('default_code', '=', 'T100191954')])
        self.assertFalse(first.yr_line)
        wizard = self._import(rows)
        self.assertEqual(Template.search_count([('default_code', 'in', ['T100191452', 'T100191954'])]), 2)
        for tmpl in first | second:
            self.assertEqual((tmpl.yr_line, tmpl.yr_brand, tmpl.yr_status), ('CN3', 'YVES ROCHER', 'Actif'))
            self.assertEqual(tmpl.yr_sub_axe_id.name, 'SOIN CHEVEUX')
            self.assertEqual(tmpl.yr_axe_id.name, 'CAPILLAIRE TEST')
        self.assertEqual((first.barcode, second.barcode), ('9990005648712', False))
        self.assertIn('Mis à jour : <b>2</b>', wizard.result)

    def test_crm_import_screen_does_the_fastmag_import(self):
        rows = [ARTICLE_HEADER, article('T100102999', '9990005911999', 'Stick', 'MAQUILLAGE TEST', 'TEINT TEST', 9)]
        crm = self.env['loyalty.product.import'].create({'file': xlsx(rows)})       # no file name sent
        crm.action_import()
        tmpl = self.env['product.template'].search([('default_code', '=', 'T100102999')])
        self.assertEqual((tmpl.yr_line, tmpl.yr_sub_axe_id.name, tmpl.yr_brand), ('CN3', 'TEINT TEST', 'YVES ROCHER'))
        self.assertIn('Articles (ancien format', crm.result)
        # an ordinary article file keeps the CRM import
        other = self.env['loyalty.product.import'].create({
            'file': base64.b64encode('Code article;Désignation;Prix\nTX1;Autre;3\n'.encode()), 'filename': 'a.csv'})
        other.action_import()
        self.assertTrue(self.env['product.template'].search([('default_code', '=', 'TX1')]))

    def test_article_barcode_conflict(self):
        rows = [ARTICLE_HEADER,
                article('TA1', '9990000000001', 'Un', 'CORPS TEST', 'DOUCHE TEST', 5),
                article('TA2', '9990000000001', 'Deux', 'CORPS TEST', 'DOUCHE TEST', 6)]
        wizard = self._import(rows)
        self.assertIn('déjà utilisé', wizard.result)
        second = self.env['product.template'].search([('default_code', '=', 'TA2')])
        self.assertTrue(second and not second.barcode)

    # ── customers ────────────────────────────────────────────────────────
    def _new_row(self, card, last, first, birthday, mobile=None, amount=None, shop=None, city='MANOUBA'):
        return [card, last, first, None, '4 RUE EL KAADINE', None, None, city, 1006, None, birthday,
                datetime(2017, 5, 11), shop, 68, datetime(2018, 3, 22), mobile, None, amount]

    def test_customers_new_file(self):
        seller = self.env['hr.employee'].create({'name': 'Vendeuse 68', 'yr_logistic_code': '68'})
        rows = [NEW_HEADER,
                self._new_row(990065941, 'KSOURI', 'AMEL', datetime(1987, 5, 15), 98461465, 43.7, 'YR_YOUG'),
                self._new_row(990070668, 'BEN', 'AIED AMEL', '20/11/0000', 24379898, 43),
                self._new_row(990065894, 'BARKA', 'MAROUA', '00/01/0000', 55140540, 32, 'YR_INCONNU')]
        wizard = self._import(rows)
        self.assertEqual(wizard.file_type, 'client_new')
        Partner = self.env['res.partner']
        amel = Partner.search([('customer_code', '=', '990065941')])
        self.assertEqual(amel.name, 'KSOURI AMEL')
        self.assertEqual(amel.barcode, '990065941')                # the old card still works at the till
        self.assertTrue(amel.crm_is_customer)
        self.assertEqual((amel.birth_day, amel.birth_month, amel.birth_year), ('15', '5', '1987'))
        self.assertEqual((amel.mobile, amel.zip, amel.city), ('98461465', '1006', 'MANOUBA'))
        self.assertEqual(amel.crm_store_id, self.store)
        self.assertEqual((amel.yr_entry_date, amel.yr_first_purchase_legacy), (date(2017, 5, 11), date(2018, 3, 22)))
        self.assertEqual(amel.sudo().yr_legacy_amount, 43.7)
        self.assertEqual(amel.yr_seller_code, '68')
        self.assertEqual(amel.yr_seller_id, seller)                  # « Créé par » = the seller of the file
        self.assertEqual(amel.country_id, self.env.ref('base.tn'))
        ben = Partner.search([('customer_code', '=', '990070668')])
        self.assertEqual((ben.name, ben.birth_day, ben.birth_month, ben.birth_year),
                         ('BEN AIED AMEL', '20', '11', False))     # 0000: day and month only
        maroua = Partner.search([('customer_code', '=', '990065894')])
        self.assertFalse(maroua.birth_day)                          # 00/01/0000: no birthday
        self.assertIn('YR_INCONNU', wizard.result)
        # the frozen migration amount is kept on re-import, empty fields are completed
        rows[1] = self._new_row(990065941, 'KSOURI', 'AMEL', datetime(1987, 5, 15), 98461465, 999, 'YR_YOUG')
        rows[1][16] = 'amel@example.com'
        self._import(rows)
        self.assertEqual(Partner.search_count([('customer_code', '=', '990065941')]), 1)
        self.assertEqual(amel.sudo().yr_legacy_amount, 43.7)
        self.assertEqual(amel.email, 'amel@example.com')
        with self.assertRaises(UserError):
            amel.sudo().write({'yr_legacy_amount': 1})

    def test_delete_imported_customers_and_sales(self):
        """Paramètres button: imported customers and sales deleted (a customer used elsewhere is archived);
        stores and sellers kept."""
        self._import([NEW_HEADER, self._new_row(990022222, 'DEL', 'ONE', datetime(1990, 1, 1), 98000002),
                      self._new_row(990033333, 'DEL', 'TWO', datetime(1990, 1, 1), 98000003)])
        Partner = self.env['res.partner'].with_context(active_test=False)
        used = Partner.search([('customer_code', '=', '990033333')])
        self.env['sale.order'].create({'partner_id': used.id})
        self.env['yr.legacy.sale.line'].create({'shop_code': 'YR_YOUG', 'ticket': 'DEL1', 'date': date.today()})
        seller = self.env['hr.employee'].create({'name': 'Gardée', 'yr_logistic_code': '9911'})
        self.env['res.config.settings'].create({}).action_yr_delete_imported()
        self.assertFalse(Partner.search([('customer_code', '=', '990022222')]))
        self.assertFalse(used.active)                                   # used by a sale order: archived
        self.assertFalse(self.env['yr.legacy.sale.line'].search_count([]))
        self.assertTrue(self.store.exists() and seller.exists())

    def test_customers_street2(self):
        """Rue 2 (and 3) → street2; customers imported before with everything in the street are fixed on re-import."""
        row = self._new_row(990011111, 'RUE', 'DEUX', datetime(1990, 1, 1), 98000001)
        row[5], row[6] = 'RESIDENCE LES PINS', 'BLOC B'
        self._import([NEW_HEADER, row])
        partner = self.env['res.partner'].search([('customer_code', '=', '990011111')])
        self.assertEqual((partner.street, partner.street2), ('4 RUE EL KAADINE', 'RESIDENCE LES PINS BLOC B'))
        partner.write({'street': '4 RUE EL KAADINE RESIDENCE LES PINS BLOC B', 'street2': False})   # old import
        self._import([NEW_HEADER, row])
        self.assertEqual((partner.street, partner.street2), ('4 RUE EL KAADINE', 'RESIDENCE LES PINS BLOC B'))

    def test_customers_civility(self):
        """Civilité / ancienne civilité / sexe: from the list (abbreviations understood, new values added);
        a later file with the card and the civility only completes the customers."""
        header = ['CardNumbre', 'SurName', 'FirstName', 'Civilité', 'Ancienne civilité', 'Sexe']
        self._import([header, [990055001, 'CIV', 'UN', 'Mme', 'MLLE', None],
                      [990055002, 'CIV', 'DEUX', 'Maître', None, 'H'],
                      [990055003, 'CIV', 'TROIS', None, None, None]])
        Partner = self.env['res.partner']
        one, two, three = (Partner.search([('customer_code', '=', c)]) for c in ('990055001', '990055002',
                                                                                  '990055003'))
        self.assertEqual(one.title, self.env.ref('base.res_partner_title_madam'))
        self.assertEqual(one.crm_old_title_id, self.env.ref('base.res_partner_title_miss'))
        self.assertEqual(one.crm_gender, 'f')                         # deduced from « Madame »
        self.assertEqual((two.title.name, two.crm_gender), ('Maître', 'h'))   # new civility added to the list
        self.assertFalse(three.title)                                # empty for now
        self._import([['CardNumbre', 'Civilite'], [990055003, 'M.']])
        self.assertEqual((three.title, three.crm_gender, three.name),
                         (self.env.ref('base.res_partner_title_mister'), 'h', 'CIV TROIS'))
        arch = Partner.get_view(self.env.ref('base.view_partner_form').id, 'form')['arch']
        self.assertIn('name="crm_old_title_id"', arch)

    def test_customers_old_file(self):
        def old(card, birthday, last, first, tel, gender=None, npai=0):
            row = [None] * len(OLD_HEADER)
            row[0], row[1], row[3], row[4], row[5], row[7], row[8], row[10] = \
                1, card, birthday, last, first, '8 RUE AHMED', '2083', 'CITE LA GAZELLE'
            row[13], row[14], row[15], row[22], row[24], row[31] = npai, 31, datetime(2014, 12, 12), 'PA', tel, gender
            return row
        rows = [OLD_HEADER,
                old('990001529', datetime(2000, 1, 23), 'AMRI ', 'BASSIMA', '22.51.36.90', 'F', 1),
                old('990001552', datetime(1978, 3, 25), 'MOATEMRI', 'CHIRAZ', '71.23.45.67')]
        wizard = self._import(rows)
        self.assertEqual(wizard.file_type, 'client_old')
        Partner = self.env['res.partner']
        bassima = Partner.search([('customer_code', '=', '990001529')])
        self.assertEqual((bassima.name, bassima.mobile, bassima.phone), ('AMRI BASSIMA', '22513690', False))
        self.assertEqual((bassima.birth_day, bassima.birth_month, bassima.birth_year), ('23', '1', False))
        self.assertEqual((bassima.crm_gender, bassima.yr_npai, bassima.yr_legacy_segment), ('f', True, 'PA'))
        chiraz = Partner.search([('customer_code', '=', '990001552')])
        self.assertEqual((chiraz.birth_year, chiraz.phone), ('1978', '71234567'))
        # year 2000 kept when the option is off
        self.env['res.partner'].search([('customer_code', '=', '990001529')]).write(
            {'birth_day': False, 'birth_month': False})
        self._import(rows, year_2000_unknown=False)
        self.assertEqual(bassima.birth_year, '2000')

    def test_turnover_file(self):
        """The fastmag purchases count in Fidélité & CRM: status, dates, turnover, last store."""
        today = date.today()
        recent, older = today - timedelta(days=20), today - timedelta(days=300)
        rows = [CA_HEADER,
                ['YR_YOUG', datetime.combine(recent, datetime.min.time()), 6, 990025739, 'FERCHICHI', 'SAMEH', 246.3],
                ['YR_XQZW', datetime.combine(older, datetime.min.time()), 6, 990025739, 'FERCHICHI', 'SAMEH', 1328.8],
                ['YR_MANAR', datetime(2026, 3, 19), 29, 999999999, 'INCONNUE', 'X', 21.8]]
        # history first: kept until the customer exists
        wizard = self._import(rows)
        self.assertEqual(wizard.file_type, 'client_ca')
        self.assertIn('Cartes sans fiche client : <b>2</b>', wizard.result)
        self.assertIn('YR_XQZW', wizard.result)                  # store codes to map
        Legacy = self.env['yr.legacy.purchase']
        self.assertEqual(Legacy.search_count([('card', '=', '990025739')]), 2)
        self._import([NEW_HEADER, self._new_row(990025739, 'FERCHICHI', 'SAMEH', datetime(1990, 1, 2))])
        sameh = self.env['res.partner'].search([('customer_code', '=', '990025739')])
        self.assertEqual(len(sameh.yr_legacy_purchase_ids), 2)
        self.assertAlmostEqual(sameh.crm_amount_total, 1575.1)
        self.assertEqual(sameh.crm_order_count, 2)
        self.assertEqual(sameh.crm_first_order_date, date(2018, 3, 22))   # first purchase of the customer file
        self.assertEqual(sameh.crm_last_order_date, recent)
        self.assertEqual(sameh.crm_last_store_id, self.store)
        self.assertEqual(sameh.crm_lifecycle, 'active')
        orders = self.env['loyalty.crm.order'].search([('partner_id', '=', sameh.id)])
        self.assertEqual(set(orders.mapped('source')), {'fastmag'})
        # re-import: amounts updated, no duplicate
        rows[1][6] = 250.0
        self._import(rows)
        self.assertEqual(Legacy.search_count([('card', '=', '990025739')]), 2)
        self.assertAlmostEqual(sameh.crm_amount_total, 1578.8)

    def test_customer_without_history_status(self):
        self._import([NEW_HEADER, self._new_row(990000001, 'ANCIENNE', 'CLIENTE', datetime(1980, 2, 3))])
        old = self.env['res.partner'].search([('customer_code', '=', '990000001')])
        self.assertEqual(old.crm_first_order_date, date(2018, 3, 22))
        self.assertEqual(old.crm_lifecycle, 'long_inactive')            # bought, long ago: not a prospect

    def test_unknown_file(self):
        wizard = self.env['yr.migration.import'].create({'file': xlsx([['a', 'b'], [1, 2]]), 'filename': 'x.xlsx'})
        self.assertFalse(wizard.file_type)
        with self.assertRaises(UserError):
            wizard.action_import()

    # ── store list ───────────────────────────────────────────────────────
    def test_store_list(self):
        Store = self.env['loyalty.store']
        azur = Store.create({'name': 'Azurq', 'code': 'AZRQ'})
        wizard = self._import([['Magasin', 'Code'], ['AZURQ', 923], ['ZEPHYRQ TEST', 920], ['eshop', 926]])
        self.assertEqual(wizard.file_type, 'stores')
        self.assertEqual((azur.code, azur.yr_fastmag_code), ('923', 'AZRQ'))       # the former code stays known
        zephyr = Store.search([('code', '=', '920')])
        self.assertEqual((zephyr.name, zephyr.store_type), ('ZEPHYRQ TEST', 'store'))
        self.assertEqual(Store.search([('code', '=', '926')]).store_type, 'website')
        # one company per store: a branch of the main company; the website's branch is flagged « Site web »
        root = self.env.company.root_id
        imported = Store.search([('code', 'in', ['923', '920', '926'])])
        self.assertEqual(len(imported.company_id), 3)
        self.assertEqual(imported.company_id.parent_id, root)
        self.assertEqual(zephyr.company_id.name, 'ZEPHYRQ TEST')
        self.assertTrue(Store.search([('code', '=', '926')]).company_id.crm_is_website)
        self.assertTrue(imported.company_id <= self.env.user.company_ids)
        # again, and without header row: nothing duplicated
        wizard = self._import([['ZEPHYRQ TEST', 920], ['NABEUL TEST', 919]])
        self.assertEqual(wizard.file_type, 'stores')
        self.assertEqual(Store.search_count([('name', '=', 'ZEPHYRQ TEST')]), 1)
        self.assertEqual(self.env['res.company'].search_count([('name', '=', 'ZEPHYRQ TEST')]), 1)  # not twice
        self.assertTrue(Store.search([('code', '=', '919'), ('name', '=', 'NABEUL TEST')]))
        # « YR_ZEPHYR » of the fastmag files is recognised from the name, and remembered
        stores = self.env['yr.migration.import']._stores()
        self.assertEqual(stores.get('YR_ZEPHYRQ'), zephyr)
        self.assertIn('YR_ZEPHYRQ', zephyr.yr_fastmag_code)

    def test_store_pos_and_sales_by_till(self):
        """Store list: a till (and a warehouse) in each store's subsidiary; the detailed sales go to that till;
        the import is recorded in the history."""
        wizard = self._import([['magasin', 'code', 'statut'], ['TILLSTOREV', 971, None], ['CLOSEDV', 972, 'fermé'],
                               ['eshopv', 973, None]])
        Store = self.env['loyalty.store'].with_context(active_test=False)
        store = Store.search([('code', '=', '971')])
        self.assertEqual(len(store.pos_config_ids), 1)
        self.assertEqual(store.pos_config_ids.company_id, store.company_id)
        self.assertFalse(Store.search([('code', '=', '972')]).pos_config_ids)      # closed: no till
        self.assertFalse(Store.search([('code', '=', '973')]).pos_config_ids)      # website: no till
        self.assertEqual((wizard.log_id.state, wizard.log_id.name), ('done', 'fichier.xlsx'))
        store.yr_fastmag_code = 'YR_TILLV'
        header = ['CodeMag', 'Utilisateur', 'Vendeur', 'Date', 'Client', 'vente', 'Barcode', 'Quantite', 'Prix',
                  'remise', 'total']
        self._import([header, ['YR_TILLV', 'RIM_T', 'SANAV', '27/11/2017', None, 1, 100191290, 1, 13, 0, 13]])
        line = self.env['yr.legacy.sale.line'].search([('shop_code', '=', 'YR_TILLV')])
        self.assertEqual(line.pos_config_id, store.pos_config_ids)

    def test_resume_interrupted_import(self):
        """An import stopped half-way: importing the same file again skips what was saved, and says so."""
        header = ['CodeMag', 'Utilisateur', 'Vendeur', 'Date', 'Client', 'vente', 'Barcode', 'Quantite', 'Prix',
                  'remise', 'total']
        rows = [header] + [['YR_YOUG', 'RIM_T', 'SANAR', '27/11/2017', None, n, 100191290, 1, 13, 0, 13]
                           for n in range(1, 6)]
        content = xlsx(rows)
        Line = self.env['yr.legacy.sale.line']
        # the first 2 tickets were saved, then the import stopped
        Line.create([{'shop_code': 'YR_YOUG', 'ticket': str(n), 'sequence': 1, 'date': date(2017, 11, 27)}
                     for n in (1, 2)])
        wizard = self.env['yr.migration.import'].create({'file': content, 'filename': 'ventes.xlsx'})
        self.env['yr.migration.log'].create({'name': 'ventes.xlsx', 'file_hash': wizard._file_hash(),
                                             'state': 'failed', 'done': 2, 'total': 5})
        self.assertIn('interrompu', wizard.previous_import)
        wizard.action_import()
        self.assertIn('déjà importées', wizard.result)
        self.assertEqual(Line.search_count([('shop_code', '=', 'YR_YOUG'), ('date', '=', date(2017, 11, 27))]), 5)
        self.assertEqual(wizard.log_id.state, 'done')
        self.assertEqual(wizard.get_progress()['state'], 'done')
        # once done: the screen says the file was imported
        again = self.env['yr.migration.import'].create({'file': content, 'filename': 'ventes.xlsx'})
        self.assertIn('déjà été importé', again.previous_import)

    def test_store_list_again_with_main_company_selected(self):
        """Stores already in their subsidiary, the session on the main company only (a normal user): re-importing a
        store list finds them (no « code unique » error, nothing duplicated)."""
        admin = self.env.ref('base.user_admin')
        root = self.env.company.root_id
        admin.company_ids |= root
        rows = [['Magasin', 'Code'], ['CARREFOURW', 981], ['GEANTW', 982]]
        self._import(rows)
        stores = self.env['loyalty.store'].search([('code', 'in', ['981', '982'])])
        self.assertEqual(len(stores.company_id), 2)
        admin.company_ids |= stores.company_id
        Wizard = self.env['yr.migration.import'].with_user(admin).with_context(allowed_company_ids=[root.id])
        for file_rows in (rows, [['Magasin ', 'Code -Magasin ', 'Num-Magasin '], ['CARREFOURW', 'YR_CARRW', 981],
                                 ['GEANTW', 'YR_GEANTW', 982]]):
            wizard = Wizard.create({'file': xlsx(file_rows), 'filename': 'magasins.xlsx'})
            wizard.action_import()
            self.assertNotIn('en erreur', wizard.result)
        self.assertEqual(self.env['loyalty.store'].search_count([('code', 'in', ['981', '982'])]), 2)
        self.assertIn('YR_CARRW', stores.filtered(lambda s: s.code == '981').yr_fastmag_code)

    def test_workbook_stores_and_sellers(self):
        """The « vendeurs / magasins » workbook: Feuil1 sellers per store, Feuil2 magasin | YR_ code | number."""
        Store = self.env['loyalty.store']
        carr = Store.create({'name': 'CARREFOURQ', 'code': '992'})
        config = self.env['pos.config'].create({'name': 'Caisse Carrefourq', 'crm_store_id': carr.id})
        sellers = [['Code magasin ', 'Vendeurs', 'Code Logistique'],
                   ['YR_CARRQ', 'RIMQ', 801], [None, 'SANAQ TEST', 802],
                   ['YR_MALLSOUQ', 'IMENQ', 803], [None, 'SANAQ TEST', 804]]
        stores = [['Magasin ', 'Code -Magasin ', 'Num-Magasin '],
                  ['CARREFOURQ', 'YR_CARRQ', 992], ['MALL OF SOUSSEQ', 'YR_MALLSOUQ', 994]]
        wizard = self.env['yr.migration.import'].create({
            'file': xlsx_sheets({'Feuil1': sellers, 'Feuil2': stores}), 'filename': 'vendeurmagasin.xlsx'})
        self.assertEqual(wizard.file_type, 'workbook')
        wizard.action_import()
        self.assertIn('Feuil2', wizard.result)
        # stores: the YR_ code of the sales files is kept on the store, the number is its code
        self.assertIn('YR_CARRQ', carr.yr_fastmag_code)
        mall = Store.search([('code', '=', '994')])
        self.assertEqual((mall.name, mall.yr_fastmag_code), ('MALL OF SOUSSEQ', 'YR_MALLSOUQ'))
        self.assertEqual(self.env['yr.migration.import']._stores().get('YR_MALLSOUQ'), mall)
        # sellers: employees of their store, with the logistic code, proposed at the store's till
        Employee = self.env['hr.employee']
        rim = Employee.search([('yr_logistic_code', '=', '801')])
        self.assertEqual((rim.name, rim.yr_store_id), ('RIMQ', carr))
        self.assertIn(rim, config.yr_seller_employee_ids)
        self.assertEqual(Employee.search([('yr_logistic_code', '=', '803')]).yr_store_id, mall)  # code carried down
        # store codes with a typing error or abbreviated, and an unknown store (created)
        mahs = Store.create({'name': 'KAIROUANQZ', 'code': '995'})
        sfax = Store.create({'name': 'BIZERTEQ NORD', 'code': '996'})
        more = [['Code magasin ', 'Vendeurs', 'Code Logistique'],
                ['YR_KAYROUANQZ', 'HAMIDAQ', 811], ['YR_BNORD', 'OLFAQ', 812], ['YR_QXZW', 'NESRINEQ', 813]]
        self._import(more)
        self.assertEqual(Employee.search([('yr_logistic_code', '=', '811')]).yr_store_id, mahs)
        self.assertEqual(Employee.search([('yr_logistic_code', '=', '812')]).yr_store_id, sfax)
        hb = Employee.search([('yr_logistic_code', '=', '813')]).yr_store_id
        self.assertEqual((hb.name, hb.yr_fastmag_code), ('QXZW', 'YR_QXZW'))
        # closed stores (statut « fermé »): kept archived, their sellers attached; a store already known under its
        # own YR_ code (Monastir Centre = YR_MONZ_C) is not taken for another abbreviation (MCENTER = Mahdia Centre)
        sfax_c = Store.create({'name': 'MonastirZ Centre', 'code': '997', 'yr_fastmag_code': 'YR_MONZ_C'})
        self._import([['magasin', 'code', 'statut'], ['Habib Bourguiba Zitouna', 998, 'fermé'],
                      ['MahdiaZ Centre', 999, 'Fermé']])
        closed = Store.with_context(active_test=False).search([('code', 'in', ['998', '999'])])
        self.assertEqual(closed.mapped('active'), [False, False])
        self._import([['Code magasin ', 'Vendeurs', 'Code Logistique'],
                      ['YR_HBZ', 'NADIAZ', 821], ['YR_MCENTER', 'INSAFZ', 822]])
        self.assertEqual(Employee.search([('yr_logistic_code', '=', '821')]).yr_store_id.code, '998')
        self.assertEqual(Employee.search([('yr_logistic_code', '=', '822')]).yr_store_id.code, '999')
        self.assertNotIn('YR_MCENTER', sfax_c.yr_fastmag_code)
        # same name, other logistic code: two different sellers
        self.assertEqual(len(Employee.search([('name', '=', 'SANAQ TEST')])), 2)
        # again: nothing duplicated
        wizard = self.env['yr.migration.import'].create({
            'file': xlsx_sheets({'Feuil1': sellers, 'Feuil2': stores}), 'filename': 'vendeurmagasin.xlsx'})
        wizard.action_import()
        self.assertEqual(Employee.search_count([('yr_logistic_code', 'in', ['801', '802', '803', '804'])]), 4)
        self.assertEqual(Store.search_count([('code', '=', '994')]), 1)

    def test_sales_detail_client_column(self):
        """Sales file « INES-YVESROCHER | CodeMag | … | Client | vente | … »: no card column, the client code (Y…)
        finds the customer through the « Liste client » file, even when that file is imported after the sales."""
        self._import([NEW_HEADER, self._new_row(990088074, 'CLIENTE', 'Y', datetime(1990, 1, 2))])
        partner = self.env['res.partner'].search([('customer_code', '=', '990088074')])
        header = ['INES-YVESROCHER', 'CodeMag', 'Utilisateur', 'Vendeur', 'Date', 'Client', 'vente', 'Barcode',
                  'Quantite', 'Prix', 'remise', 'PrixNet', 'total', 'total_1', 'valeurArticle', 'motif']
        rows = [header,
                ['INES-YVESROCHER', 'YR_YOUG', 'RIM_T', 'SANAQ', '27/11/2017', 'Y880741', 2, 100102550, 1, 34, 0, 34,
                 34, 67, 16.73, None],
                ['INES-YVESROCHER', 'YR_YOUG', 'RIM_T', 'SANAQ', '27/11/2017', 'Y880741', 2, 100190813, 1, 33, 0, 33,
                 33, 67, 16.239, None],
                ['INES-YVESROCHER', 'YR_YOUG', 'RIM_T', 'RIMQ', '27/11/2017', None, 1, 100191290, 1, 13, 0, 13, 13, 13,
                 6.397, None],
                ['INES-YVESROCHER', 'YR_YOUG', 'RIM_T', 'MARWAQ', '27/11/2017', 'Y880741', 25, 100195183, 1, 8, 100, 0,
                 0, 132, 3.937, 'CADEAU_PTS']]
        wizard = self._import(rows)
        self.assertEqual(wizard.file_type, 'sales_detail')
        self.assertIn('code client sans carte connue', wizard.result)
        lines = self.env['yr.legacy.sale.line'].search([('shop_code', '=', 'YR_YOUG'), ('date', '=', date(2017, 11, 27))])
        self.assertEqual(len(lines), 4)
        self.assertEqual(lines.mapped('store_id'), self.store)
        self.assertFalse(lines.partner_id)                           # code not known yet
        self.assertEqual(set(lines.mapped('legacy_client_id')), {'Y880741', False})
        self.assertEqual(lines.filtered(lambda l: l.ticket == '25').reason, 'CADEAU_PTS')
        # fastmag names kept as text even without any employee
        self.assertEqual(set(lines.mapped('seller_name')), {'SANAQ', 'RIMQ', 'MARWAQ'})
        self.assertEqual(set(lines.mapped('cashier_name')), {'RIM_T'})
        self.assertFalse(lines.seller_id)                            # no employee guessed
        # « Liste client »: the sales imported before find their customer
        codes = self._import([CODES_HEADER, ['Y880741', 990088074, 'CLIENTE Y', datetime(2017, 1, 1), 3, 'CLYVRC']])
        self.assertEqual(codes.file_type, 'client_codes')
        self.assertEqual(lines.filtered(lambda l: l.ticket == '2').partner_id, partner)
        self.assertFalse(lines.filtered(lambda l: l.ticket == '1').partner_id)
        self.assertEqual(partner.yr_client_code, 'Y880741')
        purchase = self.env['yr.legacy.purchase'].search([('card', '=', '990088074')])
        self.assertAlmostEqual(sum(purchase.mapped('amount')), 67.0)
        # sales imported after the codes: found directly
        rows[1][6] = rows[2][6] = 3
        wizard = self._import(rows)
        self.assertIn('rattachées au client par le code Y', wizard.result)

    def test_client_codes_and_customers_without_card(self):
        """No real card (empty, 0, 000000, « numCliente »): not imported as a customer, listed in the report."""
        rows = [NEW_HEADER, self._new_row(990077001, 'AVEC', 'CARTE', datetime(1990, 1, 2)),
                self._new_row(0, 'SANS', 'CARTE', datetime(1990, 1, 2)),
                self._new_row('000000', 'ZERO', 'CARTE', datetime(1990, 1, 2))]
        rows.append(self._new_row(None, 'VIDE', 'CARTE', datetime(1990, 1, 2)))
        wizard = self._import(rows)
        Partner = self.env['res.partner'].with_context(active_test=False)
        self.assertTrue(Partner.search([('customer_code', '=', '990077001')]))
        self.assertFalse(Partner.search([('name', 'in', ['SANS CARTE', 'ZERO CARTE', 'VIDE CARTE'])]))
        self.assertIn('non importés) : <b>3</b>', wizard.result)
        report = self._report(wizard)
        self.assertEqual(sum(1 for row in report['Détail'] if row[3] == 'Non importé'), 3)
        self.assertTrue(any(row[1] == 'Clients sans vrai n° de carte fidélité (non importés)'
                            for row in report['Résumé']))
        # « Liste client »: codes kept, only real cards linked to customers
        wizard = self._import([CODES_HEADER, ['Y1', 'numCliente', 'nom prenom', datetime(2017, 1, 1), 0, 'CLYVRC'],
                               ['Y77001', 990077001, 'AVEC CARTE', datetime(2017, 1, 1), 4, 'CLYVRC'],
                               ['Y64', '000000', 'BEN AHMED', datetime(2014, 8, 8), 0, 'CLYVRC']])
        Code = self.env['yr.client.code']
        self.assertEqual(Code.search([('code', '=', 'Y77001')]).partner_id.customer_code, '990077001')
        self.assertFalse(Code.search([('code', '=', 'Y64')]).card)
        self.assertEqual(Code.search_count([('code', 'in', ['Y1', 'Y64', 'Y77001'])]), 3)
        self.assertFalse(Partner.search([('name', '=', 'BEN AHMED')]))
        self.assertIn('Codes sans vraie carte (ventes anonymes) : <b>2</b>', wizard.result)
        # a customer file imported after the codes gets its code; the CA file finds the card from the client number
        self._import([NEW_HEADER, self._new_row(990077002, 'APRES', 'CODES', datetime(1990, 1, 2))])
        self._import([CODES_HEADER, ['Y77002', 990077002, 'APRES CODES', datetime(2017, 1, 1), 0, 'CLYVRC']])
        after = Partner.search([('customer_code', '=', '990077002')])
        self.assertEqual(after.yr_client_code, 'Y77002')
        self._import([CA_HEADER, ['YR_YOUG', datetime(2025, 8, 26), 77002, None, 'APRES', 'CODES', 50.0]])
        self.assertEqual(after.yr_legacy_purchase_ids.amount, 50.0)

    def _report(self, wizard):
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(wizard.log_id.report_file)))
        return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}

    # ── sales detail ─────────────────────────────────────────────────────
    def _detail(self, ticket, day, cashier, seller, card, ref, ean, name, qty, price, discount, total, reason=None):
        return ['YR_YOUG', ticket, day, cashier, seller, 168479 if card else 0, card, 'SOUISSI' if card else None,
                'WAEL' if card else None, ref, ean, 56748, name, 'MAQUILLAGE', 'CN3', qty, price, discount, total,
                reason, None]

    def test_sales_detail(self):
        product = self.env['product.product'].create({'name': 'Crayon khôl test', 'default_code': '900106131',
                                                      'barcode': '3660009567488', 'list_price': 23})
        cashier = self.env['hr.employee'].create({'name': 'Rim Caissière', 'yr_fastmag_name': 'RIM'})
        self._import([NEW_HEADER, self._new_row(9940133181, 'SOUISSI', 'WAEL', datetime(1990, 1, 2))])
        day = datetime.combine(date.today() - timedelta(days=3), datetime.min.time())
        rows = [DETAIL_HEADER,
                self._detail(9199659, day, 'RIM', 'JBELI TEST', None, 900106131, 3660009567488, 'CRAYON', 1, 23, 0, 23),
                self._detail(9199659, day, 'RIM', 'JBELI TEST', None, 'DTT', 75229, 'DROIT DE TIMBRE', 1, 0.1, 0, 0.1),
                self._detail(9199667, day, 'RIM', 'SLIMENI TEST', 9940133181, 'X1', 3660009567488, 'CRAYON', 2, 23, 0,
                             46),
                self._detail(9199667, day, 'RIM', 'SLIMENI TEST', 9940133181, 'INCONNU', None, 'CREME', 1, 18, 100, 0,
                             'CADEAU_PTS'),
                self._detail(9199667, day, 'RIM', 'SLIMENI TEST', 9940133181, 'DTT', 75229, 'DROIT DE TIMBRE', 1, 0.1,
                             0, 0.1)]
        wizard = self._import(rows, create_sellers=True)
        self.assertEqual(wizard.file_type, 'sales_detail')
        Line = self.env['yr.legacy.sale.line']
        lines = Line.search([('shop_code', '=', 'YR_YOUG'), ('ticket', 'in', ['9199659', '9199667'])])
        self.assertEqual(len(lines), 5)
        first = lines.filtered(lambda l: l.ticket == '9199659' and l.sequence == 1)
        self.assertEqual((first.product_id, first.store_id, first.cashier_id), (product, self.store, cashier))
        self.assertEqual((first.seller_name, first.seller_id.name), ('JBELI TEST', 'JBELI TEST'))  # seller created
        self.assertTrue(lines.filtered(lambda l: l.product_ref == 'DTT').mapped('is_stamp') == [True, True])
        self.assertEqual(lines.filtered(lambda l: l.product_ref == 'X1').product_id, product)       # by EAN
        gift = lines.filtered(lambda l: l.reason == 'CADEAU_PTS')
        self.assertEqual((gift.discount, gift.amount, gift.product_id.id), (100, 0, False))
        wael = self.env['res.partner'].search([('customer_code', '=', '9940133181')])
        self.assertEqual(lines.filtered(lambda l: l.ticket == '9199667').partner_id, wael)
        # purchase history of the customer: one visit, the total of the ticket
        purchase = self.env['yr.legacy.purchase'].search([('card', '=', '9940133181')])
        self.assertEqual((len(purchase), purchase.store_id), (1, self.store))
        self.assertAlmostEqual(purchase.amount, 46.1)
        self.assertEqual(wael.crm_last_order_date, day.date())
        # re-import: nothing duplicated, employees not created twice
        self._import(rows)
        self.assertEqual(Line.search_count([('shop_code', '=', 'YR_YOUG'), ('ticket', 'in', ['9199659', '9199667'])]), 5)
        self.assertEqual(self.env['hr.employee'].search_count([('name', '=', 'JBELI TEST')]), 1)
        self.assertEqual(self.env['yr.legacy.purchase'].search_count([('card', '=', '9940133181')]), 1)
        # a seller created later in Employés is attached with « Rattacher aux employés »
        self._import(rows, create_sellers=False)
        slimeni = lines.filtered(lambda l: l.seller_name == 'SLIMENI TEST')
        slimeni.seller_id.unlink()
        self.assertFalse(slimeni.seller_id)
        employee = self.env['hr.employee'].create({'name': 'Slimeni Maram', 'yr_fastmag_name': 'slimeni test'})
        Line._yr_make_pos_orders()
        Line._link_employees()
        self.assertEqual(slimeni.seller_id, employee)
        self.assertEqual(slimeni.pos_order_id.yr_seller_id, employee)   # the POS order follows

    # ── new files: client detail, Sage X3 articles, coupons, motifs ──────
    def test_client_detail_and_sage_articles(self):
        header = ['Client', 'Nom', 'Prenom', 'Civilite', 'Adresse1', 'codepostal', 'Ville', 'Pays', 'Telephone',
                  'Portable', 'Email', 'JourNaissance', 'MoisNaissance', 'Sexe', 'CarteFidelite', 'MagasinCreation',
                  'DateCreation', 'totalca', 'PubParEMail', 'PubParSMS', 'AnneeNaissance']
        rows = [header,
                [1, 'nom', 'prenom', None, 'adresse', 'code postal', 'ville', None, 'telephone', 'gsm', 'e-mail', 0, 0,
                 'F', 'numCliente', 'CENTRAL', datetime(2017, 1, 1), 0, 1, 1, 0],
                [77001, 'DETAIL', 'NADIA', ' MME', '7 RUE TEST', 1000, 'TUNIS', None, None, 99230596, 'N@T.TN', 14, 3,
                 'F', 9940111111, 'CENTRAL', datetime(2012, 8, 22), 0, 1, 0, 1985],
                [77002, 'DETAIL', 'SAMIA', ' 7 RUE MOHAM', '8 RUE TEST', 1000, 'TUNIS', 'France', None, 96406255, None,
                 0, 0, 'F', 9940222222, 'CENTRAL', datetime(2017, 7, 19), 0, 0, 0, 0]]
        wizard = self._import(rows)
        self.assertEqual(wizard.file_type, 'client_detail')
        Partner = self.env['res.partner']
        nadia = Partner.search([('customer_code', '=', '9940111111')])
        self.assertEqual((nadia.yr_client_code, nadia.title.shortcut, nadia.crm_gender), ('Y77001', 'Mrs.', 'f'))
        self.assertEqual((nadia.birth_day, nadia.birth_month, nadia.birth_year), ('14', '3', '1985'))
        self.assertEqual((nadia.mobile, nadia.optin_email, nadia.optin_sms), ('99230596', True, False))
        self.assertFalse(Partner.search([('customer_code', '=', 'numCliente')]))
        samia = Partner.search([('customer_code', '=', '9940222222')])
        self.assertFalse(samia.title)                   # an address in the civility column is not a title
        self.assertEqual(samia.country_id.code, 'FR')
        self.assertNotIn('CENTRAL', wizard.result)      # CENTRAL is not a store

        sage = ['Article', 'Catégori', 'Désignation    ', 'Code à barre', 'Etat', 'Marque ', 'Marque', 'Ligne', 'Axe',
                'Sous Axe', 'Nature', 'CREDAT_0']
        rows = [sage,
                ['100000000', 'CAFV', "BON D'ACHAT", '100000000', 1, 'M0074', 'YVES ROCHER', 'DIVERS', 'DIVERS',
                 'DIVERS', '0', datetime(2023, 11, 30)],
                ['SAGE100', 'CDISC', 'LUM 4 VAO BLANC', '5689500', 1, 'M0074', 'YVES ROCHER', 'LUM 4', 'MAQ SAGE',
                 'ONGLES SAGE', 'VERNIS', datetime(2014, 5, 28)],
                ['SAGE101', 'CDISC', 'LUM 4 VAO ROSE', '5712100', 0, 'M0074', 'YVES ROCHER', 'LUM 4', 'MAQ SAGE',
                 'ONGLES SAGE', 'VERNIS', datetime(2014, 5, 28)]]
        wizard = self._import(rows, archive_inactive=True)
        self.assertEqual(wizard.file_type, 'article_sage')
        Template = self.env['product.template'].with_context(active_test=False)
        self.assertFalse(Template.search([('default_code', '=', '100000000')]))      # the coupon article
        blanc = Template.search([('default_code', '=', 'SAGE100')])
        self.assertEqual((blanc.name, blanc.barcode, blanc.yr_brand, blanc.yr_axe_id.name, blanc.yr_sub_axe_id.name),
                         ('LUM 4 VAO BLANC', '5689500', 'YVES ROCHER', 'MAQ SAGE', 'ONGLES SAGE'))
        self.assertEqual(blanc.yr_status, 'Actif')
        self.assertFalse(Template.search([('default_code', '=', 'SAGE101')]).active)

    def test_customers_without_card_are_imported_with_their_code(self):
        header = ['Client', 'Nom', 'Prenom', 'CarteFidelite', 'JourNaissance', 'MoisNaissance', 'AnneeNaissance',
                  'PubParEMail']
        rows = [header,
                [88001, 'PREMIER', 'CARTE', 9951001, 0, 0, 0, 1],
                [88002, 'DEUXIEME', 'MEME CARTE', 9951001, 0, 0, 0, 1],
                [88003, 'SANS', 'CARTE', None, 29, 2, 1995, 0],
                [None, 'SANS', 'CODE', None, 0, 0, 0, 0]]
        wizard = self._import(rows)
        self.assertEqual(wizard.file_type, 'client_detail')
        Partner = self.env['res.partner'].with_context(active_test=False)
        first = Partner.search([('yr_client_code', '=', 'Y88001')])
        second = Partner.search([('yr_client_code', '=', 'Y88002')])
        third = Partner.search([('yr_client_code', '=', 'Y88003')])
        self.assertEqual(first.customer_code, '9951001')
        self.assertFalse(second.customer_code)                          # the shared card stays on the first
        self.assertIn('9951001', second.comment or '')
        self.assertFalse(third.customer_code)                           # no card number is invented
        self.assertEqual((third.birth_day, third.birth_month, third.birth_year), ('29', '2', False))
        self.assertFalse(Partner.search([('name', 'ilike', 'SANS CODE')]))   # nothing to recognise it by
        self.assertIn('importés sans n° de carte', wizard.result)
        self.assertIn('Cartes partagées', wizard.result)
        # the sales detail finds each customer by its client code, even when the card is shared
        day = datetime.combine(date.today() - timedelta(days=4), datetime.min.time())
        line_a = self._detail(6199001, day, 'RIM', 'VEND', 9951001, 'ZC1', 3660009570001, 'CRAYON', 1, 20, 0, 20)
        line_a[5] = 88002                                                   # client code of the second customer
        line_b = self._detail(6199002, day, 'RIM', 'VEND', None, 'ZC2', 3660009570002, 'CREME', 1, 18, 0, 18)
        line_b[5] = 88003                                                   # no card at all
        self._import([DETAIL_HEADER, line_a, line_b])
        Line = self.env['yr.legacy.sale.line']
        self.assertEqual(Line.search([('ticket', '=', '6199001')]).partner_id, second)
        self.assertEqual(Line.search([('ticket', '=', '6199002')]).partner_id, third)
        # re-import: nobody duplicated
        self._import(rows)
        self.assertEqual(Partner.search_count([('yr_client_code', 'in', ['Y88001', 'Y88002', 'Y88003'])]), 3)
        # the option off: the old behaviour (customers without a card are left out)
        off = self._import([header, [88009, 'AUTRE', 'SANS CARTE', None, 0, 0, 0, 0]], import_no_card=False)
        self.assertFalse(Partner.search([('yr_client_code', '=', 'Y88009')]))
        self.assertIn('non importés) : <b>1</b>', off.result)

    def test_old_xls_file(self):
        try:
            import xlwt
        except ImportError:
            self.skipTest('xlwt missing')
        book = xlwt.Workbook()
        sheet = book.add_sheet('Sheet1')
        header = ['Article', 'Catégori', 'Désignation    ', 'Code à barre', 'Etat', 'Marque ', 'Marque', 'Ligne']
        for c, value in enumerate(header):
            sheet.write(0, c, value)
        for c, value in enumerate(['XLS100', 'CDISC', 'ARTICLE XLS', '5600001', 1, 'M1', 'YVES ROCHER', 'L1']):
            sheet.write(1, c, value)
        out = io.BytesIO()
        book.save(out)
        wizard = self.env['yr.migration.import'].create({'file': base64.b64encode(out.getvalue()),
                                                         'filename': 'Liste des article Sage x3.xls'})
        self.assertEqual(wizard.file_type, 'article_sage')
        wizard.action_import()
        self.assertEqual(self.env['product.template'].search([('default_code', '=', 'XLS100')]).yr_brand,
                         'YVES ROCHER')

    def test_several_motifs_follow_the_setting(self):
        day = datetime.combine(date.today() - timedelta(days=2), datetime.min.time())
        icp = self.env['ir.config_parameter'].sudo()
        for rule, ticket, expected in (('all', 7199010, ['CADEAU_PTS', 'ECHANGE']), ('first', 7199011, ['ECHANGE'])):
            icp.set_param('loyalty_tier_marketing.allocation_rule', rule)
            self._import([DETAIL_HEADER, self._detail(ticket, day, 'RIM', 'VEND', None, 'ZM1', 3660009561001,
                                                      'CRAYON', 1, 23, 0, 23, 'ECHANGE ; CADEAU_PTS')])
            self.env['yr.legacy.sale.line']._yr_make_pos_orders()
            order = self.env['pos.order'].search([('yr_fastmag_ticket', '=', 'YR_YOUG/%s' % ticket)])
            self.assertEqual(sorted(order.lines.crm_motif_ids.mapped('name')), expected)
            self.assertEqual(order.lines.crm_motif_id.name, 'ECHANGE')
            view = self.env['loyalty.crm.order.line'].search([('order_id.reference', 'ilike', str(ticket))])
            self.assertEqual(view.motif_names, ' ; '.join(expected))

    def test_motif_deduced_from_discount(self):
        day = datetime.combine(date.today() - timedelta(days=2), datetime.min.time())
        rows = [DETAIL_HEADER,
                self._detail(7199020, day, 'RIM', 'VEND', None, 'ZD1', 3660009562001, 'CRAYON', 1, 20, 25, 15),
                self._detail(7199020, day, 'RIM', 'VEND', None, 'ZD2', 3660009562002, 'CREME', 1, 18, 0, 18)]
        self._import(rows)
        lines = self.env['yr.legacy.sale.line'].search([('ticket', '=', '7199020')])
        self.assertEqual({l.product_ref: l.motif_id.name for l in lines}, {'ZD1': 'Remise (déduite)', 'ZD2': False})
        self._import([DETAIL_HEADER, self._detail(7199021, day, 'RIM', 'VEND', None, 'ZD3', 3660009562003, 'X', 1, 20,
                                                  25, 15)], deduce_motif=False)
        self.assertFalse(self.env['yr.legacy.sale.line'].search([('ticket', '=', '7199021')]).motif_id)

    def test_reimport_completes_and_updates_lines(self):
        day = datetime.combine(date.today() - timedelta(days=3), datetime.min.time())
        Line = self.env['yr.legacy.sale.line']
        self._import([DETAIL_HEADER,
                      self._detail(7199101, day, 'RIM', 'VEND', None, 'ZZ1', 3660009560001, 'CRAYON', 1, 23, 0, 23)])
        first = Line.search([('ticket', '=', '7199101')])
        self.assertEqual(len(first), 1)
        self.assertEqual(first.name, 'CRAYON / 3660009560001')          # unknown article: name / barcode
        # second import: a coupon in front and a line that was missing; the known line is updated, not doubled
        wizard = self._import([DETAIL_HEADER,
                               self._detail(7199101, day, 'RIM', 'VEND', None, 'Y13051967', None, 'COUPON', 1, 10, 100, 0),
                               self._detail(7199101, day, 'RIM', 'VEND', None, 'ZZ1', 3660009560001, 'CRAYON', 2, 23, 0, 46),
                               self._detail(7199101, day, 'RIM', 'VEND', None, 'ZZ2', 3660009560002, 'CREME', 1, 18, 0, 18)])
        lines = Line.search([('ticket', '=', '7199101')])
        self.assertEqual(len(lines), 2)
        self.assertIn(first, lines)
        self.assertEqual(first.qty, 2)
        self.assertEqual(len(set(lines.mapped('sequence'))), 2)
        self.assertEqual(sorted(lines.mapped('product_ref')), ['ZZ1', 'ZZ2'])
        self.assertIn('manquantes', wizard.result)

    def test_coupons_skipped_and_motifs(self):
        day = datetime.combine(date.today() - timedelta(days=2), datetime.min.time())
        rows = [DETAIL_HEADER,
                self._detail(7199001, day, 'RIM', 'VEND', None, 'ZZ1', 3660009560001, 'CRAYON', 1, 23, 0, 23,
                             'ECHANGE'),
                self._detail(7199001, day, 'RIM', 'VEND', None, 'Y13051967', None, 'COUPON', 1, 10, 100, 0),
                self._detail(7199001, day, 'RIM', 'VEND', None, '100000000', None, "BON D'ACHAT", 1, 10, 100, 0),
                self._detail(7199001, day, 'RIM', 'VEND', None, 'Y88074', 3660009560003, 'ARTICLE Y', 1, 12, 0, 12),
                self._detail(7199001, day, 'RIM', 'VEND', None, 'ZZ2', 3660009560002, 'CREME', 1, 18, 100, 0,
                             'CADEAU_PTS')]
        wizard = self._import(rows)
        lines = self.env['yr.legacy.sale.line'].search([('ticket', '=', '7199001')])
        self.assertEqual(sorted(lines.mapped('product_ref')), ['100000000', 'Y88074', 'ZZ1', 'ZZ2'])      # only the « coupon » label is left out
        self.assertIn('coupon', wizard.result.lower())
        motifs = lines.mapped('motif_id')
        self.assertEqual(sorted(motifs.mapped('name')), ['CADEAU_PTS', 'ECHANGE', 'Remise (déduite)'])   # + the free bon d'achat line
        self.assertEqual(motifs.filtered(lambda m: m.name == 'ECHANGE').type, 'exchange')
        self.env['yr.legacy.sale.line']._yr_make_pos_orders()
        order = self.env['pos.order'].search([('yr_fastmag_ticket', '=', 'YR_YOUG/7199001')])
        self.assertEqual(sorted(order.lines.mapped('crm_motif_id.name')), ['CADEAU_PTS', 'ECHANGE', 'Remise (déduite)'])
        # a campaign is a motif too: its name typed as the motif counts the ticket as a return of the campaign
        campaign = self.env['loyalty.tier.campaign'].create({'name': 'Fête des mères test', 'date_end': date.today()})
        self.assertEqual(self.env['loyalty.motif'].search([('campaign_id', '=', campaign.id)]).type, 'campaign')
        rows = [DETAIL_HEADER,
                self._detail(7199002, day, 'RIM', 'VEND', None, 'ZZ1', 3660009560001, 'CRAYON', 2, 23, 0, 46,
                             'Fête des mères test'),
                self._detail(7199002, day, 'RIM', 'VEND', None, 'ZZ2', 3660009560002, 'CREME', 1, 18, 0, 18)]
        self._import(rows)
        self.env['yr.legacy.sale.line']._yr_make_pos_orders()
        campaign.action_refresh_results()
        self.assertEqual((campaign.return_count, campaign.revenue), (1, 64.0))
        data = self.env['loyalty.crm.dashboard'].get_data(campaign_id=campaign.id, date_from=day.date())
        self.assertEqual(data['current']['tickets'], 1)
        self.assertEqual([m['name'] for m in data['motifs']], ['Fête des mères test'])

    def test_sales_become_pos_orders(self):
        """Each imported ticket becomes a POS order of a closed « historique fastmag » session; the CRM counts it
        once; re-import, delete buttons."""
        product = self.env['product.product'].create({'name': 'Crayon PdV test', 'default_code': '900106199',
                                                      'barcode': '3660009567499', 'list_price': 23})
        self._import([NEW_HEADER, self._new_row(9940133199, 'POS', 'CLIENTE', datetime(1990, 1, 2))])
        wael = self.env['res.partner'].search([('customer_code', '=', '9940133199')])
        day = datetime.combine(date.today() - timedelta(days=5), datetime.min.time())
        rows = [DETAIL_HEADER,
                self._detail(8199659, day, 'RIM', 'JBELI POS', None, 900106199, 3660009567499, 'CRAYON', 1, 23, 0, 23),
                self._detail(8199659, day, 'RIM', 'JBELI POS', None, 'DTT', 75229, 'DROIT DE TIMBRE', 1, 0.1, 0, 0.1),
                self._detail(8199667, day, 'RIM', 'SLIMENI POS', 9940133199, 'X1', 3660009567499, 'CRAYON', 2, 23, 0,
                             46),
                self._detail(8199667, day, 'RIM', 'SLIMENI POS', 9940133199, 'INCONNU9', None, 'CREME', 1, 18, 100, 0,
                             'CADEAU_PTS')]
        wizard = self._import(rows)
        self.assertIn('Point de Vente › Commandes', wizard.result)
        Line = self.env['yr.legacy.sale.line']
        Line._yr_make_pos_orders()
        Order = self.env['pos.order']
        orders = Order.search([('yr_fastmag_ticket', 'in', ['YR_YOUG/8199659', 'YR_YOUG/8199667'])])
        self.assertEqual(len(orders), 2)
        self.assertEqual(set(orders.mapped('state')), {'done'})
        session = orders.session_id
        self.assertEqual((len(session), session.state, session.yr_fastmag), (1, 'closed', True))
        self.assertEqual(session.config_id.crm_store_id, self.store)
        ticket = orders.filtered(lambda o: o.yr_fastmag_ticket == 'YR_YOUG/8199667')
        self.assertEqual((ticket.partner_id, ticket.amount_total, len(ticket.lines)), (wael, 46.0, 2))
        self.assertEqual((ticket.yr_seller_name, ticket.yr_cashier_name), ('SLIMENI POS', 'RIM'))
        self.assertEqual(ticket.date_order.date(), day.date())
        unknown = ticket.lines.filtered(lambda l: l.full_product_name.startswith('CREME'))
        self.assertEqual(unknown.product_id.default_code, 'FASTMAG-INCONNU')
        self.assertEqual(ticket.lines.filtered(lambda l: l.qty == 2).product_id, product)
        self.assertFalse(orders.picking_ids)
        # the CRM counts the ticket once (POS order, not also the fastmag history of that day)
        self.assertAlmostEqual(wael.crm_amount_total, 46.0)
        self.assertEqual(Line.search([('ticket', '=', '8199667')]).pos_order_id, ticket)
        # an article created later replaces « non retrouvé » on the order
        cream = self.env['product.product'].create({'name': 'Crème', 'default_code': 'INCONNU9'})
        self.env['res.config.settings'].create({}).action_yr_make_pos_orders()
        self.assertEqual(unknown.product_id, cream)
        # re-import: orders made again, no duplicate
        self._import(rows, update_purchases=True, create_sellers=False, archive_inactive=False)
        Line._yr_make_pos_orders()
        self.assertEqual(Order.search_count([('yr_fastmag_ticket', 'like', 'YR_YOUG/81996')]), 2)
        # « Commandes PdV seulement »: lines kept, orders made again on demand
        settings = self.env['res.config.settings'].create({})
        settings.action_yr_delete_pos_orders()
        self.assertFalse(Order.search_count([('yr_fastmag_ticket', '!=', False)]))
        self.assertEqual(Line.search_count([('ticket', 'in', ['8199659', '8199667'])]), 4)
        Line._yr_make_pos_orders()
        self.assertEqual(Order.search_count([('yr_fastmag_ticket', 'like', 'YR_YOUG/81996')]), 2)
        # « Clients seulement »: sales and orders kept, the customer comes back with the next import
        settings.action_yr_delete_customers()
        self.assertFalse(wael.exists() and wael.active)
        self.assertEqual(Order.search_count([('yr_fastmag_ticket', 'like', 'YR_YOUG/81996')]), 2)
        self._import([NEW_HEADER, self._new_row(9940133199, 'POS', 'CLIENTE', datetime(1990, 1, 2))])
        again = self.env['res.partner'].search([('customer_code', '=', '9940133199')])
        self.assertEqual(Order.search([('yr_fastmag_ticket', '=', 'YR_YOUG/8199667')]).partner_id, again)
        # « Ventes seulement »: lines and orders gone
        settings.action_yr_delete_sales()
        self.assertFalse(Line.search_count([]))
        self.assertFalse(Order.search_count([('yr_fastmag_ticket', '!=', False)]))

    # ── seller chosen at the till ────────────────────────────────────────
    def test_pos_order_seller(self):
        seller = self.env['hr.employee'].create({'name': 'Vendeuse test'})
        config = self.env['pos.config'].create({'name': 'Caisse vendeur test'})
        self.assertEqual(config.yr_seller_mode, 'none')      # enabled on the existing tills at install
        config.yr_seller_mode = 'required'
        self.assertIn({'id': seller.id, 'name': 'Vendeuse test'}, config.yr_seller_data)
        other = self.env['hr.employee'].create({'name': 'Autre test'})
        config.yr_seller_employee_ids = other
        self.assertEqual(config.yr_seller_data, [{'id': other.id, 'name': 'Autre test'}])
        config.yr_seller_mode = 'none'
        self.assertEqual(config.yr_seller_data, [{'id': other.id, 'name': 'Autre test'}])
        Order = self.env['pos.order']
        vals = Order._yr_seller_vals({'yr_seller_ref': seller.id})
        self.assertEqual(vals['yr_seller_id'], seller.id)
        self.assertFalse(Order._yr_seller_vals({'yr_seller_ref': 0})['yr_seller_id'])
        self.assertFalse(Order._yr_seller_vals({'yr_seller_ref': 999999999})['yr_seller_id'])
        self.assertEqual(Order._yr_seller_vals({'yr_seller_id': seller.id})['yr_seller_ref'], seller.id)
