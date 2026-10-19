# -*- coding: utf-8 -*-
"""Import of the fastmag files (articles, customers 2017-2026, customers of the old system, 2-year turnover).

One screen: the kind of file is recognised from its header row. Re-importing a file updates the records it
created (matched on article code / barcode and customer card number) instead of duplicating them.
"""
import base64
import difflib
import hashlib
import json
import time
import csv
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime

from odoo import _, api, fields, models
from odoo.modules import module as odoo_module
from odoo.exceptions import UserError
from odoo.tools import SQL, html_escape

BATCH = 500
FILE_TYPES = [
    ('article_sage', 'Articles (Article, Catégori, Code à barre, Axe, Sous Axe, Nature…)'),
    ('client_detail', 'Clients détail (Client, CarteFidelite, Civilite…)'),
    ('article', 'Articles (ancien format, remplacé par le fichier Articles)'),
    ('client_new', 'Clients 2017-2026 (ancien format, remplacé par « Clients détail »)'),
    ('client_old', 'Clients ancien système (ancien format, remplacé par « Clients détail »)'),
    ('client_codes', 'Codes clients fastmag (CodeClient Y…, carte fidélité)'),
    ('client_ca', 'CA clients (2 dernières années)'),
    ('sales_detail', 'Détail des ventes (ticket, utilisateur, vendeur)'),
    ('stores', 'Liste des magasins (magasin, code)'),
    ('sellers', 'Vendeurs par magasin (code magasin, vendeur, code logistique)'),
    ('workbook', 'Classeur à plusieurs feuilles (magasins, vendeurs…)'),
]
# a workbook: its sheets are imported in this order (stores first: the other files refer to them)
SHEET_ORDER = ['stores', 'sellers', 'article_sage', 'article', 'client_detail', 'client_new', 'client_old', 'client_codes', 'client_ca',
               'sales_detail']
# columns of each file (normalised header → key); the first ones identify the file
COLUMNS = {
    'article': {
        'refarticle': 'code', 'eanbar': 'barcode', 'site': 'site', 'categor': 'category', 'categorie': 'category',
        'designation': 'name', 'marque': 'brand', 'ligne': 'line', 'axe': 'axe', 'sousaxe': 'sub_axe',
        'nature': 'nature', 'cont': 'packaging', 'genre': 'gender', 'statut': 'status', 'prix': 'price',
        'un': 'uom', 'datecreati': 'create_date', 'datecreation': 'create_date', 'fourncode': 'supplier_code',
        'fourndes': 'supplier_name',
    },
    'article_sage': {
        'article': 'code', 'categori': 'category', 'designation': 'name', 'codeabarre': 'barcode', 'etat': 'status',
        'marque': 'brand', 'ligne': 'line', 'axe': 'axe', 'sousaxe': 'sub_axe', 'nature': 'nature',
        'credat0': 'create_date',
    },
    'client_detail': {
        'client': 'fm_id', 'nom': 'lastname', 'prenom': 'firstname', 'civilite': 'title', 'adresse1': 'street',
        'adresse2': 'street2', 'adresse3': 'street3', 'codepostal': 'zip', 'ville': 'city', 'pays': 'country',
        'telephone': 'phone', 'portable': 'mobile', 'email': 'email', 'journaissance': 'birth_d',
        'moisnaissance': 'birth_m', 'anneenaissance': 'birth_y', 'sexe': 'gender', 'cartefidelite': 'card',
        'magasincreation': 'shop', 'datecreation': 'entry_date', 'totalca': 'amount',
        'pubparemail': 'optin_email', 'pubparsms': 'optin_sms', 'utilisateurcreation': 'seller',
        'datedernierachat': 'last_purchase',
    },
    'client_new': {
        'cardnumbre': 'card', 'cardnumber': 'card', 'surname': 'lastname', 'firstname': 'firstname', 'id': 'fm_id',
        'adress': 'street', 'address': 'street', 'street2': 'street2', 'street3': 'street3', 'city': 'city',
        'postalcode': 'zip', 'phone': 'phone', 'birthday': 'birthday', 'entrancedate': 'entry_date',
        'codeshop': 'shop', 'codeseller': 'seller', 'firstpurchase': 'first_purchase', 'mobile': 'mobile',
        'email': 'email', 'montantttc': 'amount',
        'civilite': 'title', 'civility': 'title', 'titre': 'title', 'title': 'title',
        'anciennecivilite': 'old_title', 'oldcivility': 'old_title', 'ancienneciv': 'old_title',
        'sexe': 'gender', 'genre': 'gender', 'gender': 'gender', 'sex': 'gender',
    },
    'client_old': {
        'numcliente': 'card', 'numclienteold': 'old_code', 'naissance': 'birthday', 'nomx': 'lastname',
        'nom': 'lastname', 'prenom': 'firstname', 'rue1': 'street', 'rue2': 'street2', 'rue3': 'street3',
        'cp': 'zip', 'ville': 'city', 'npai': 'npai', 'numvendeuse': 'seller', 'segment': 'segment',
        'tel2': 'phone', 'sexe': 'gender', 'dateexcel': 'entry_date',
        'civilite': 'title', 'civility': 'title', 'titre': 'title',
        'anciennecivilite': 'old_title', 'oldcivility': 'old_title',
    },
    'client_codes': {
        'codeclient': 'code', 'cartefidelite': 'card', 'nom': 'name', 'datecreation': 'date',
        'fidelitepoint': 'points', 'fidelitepoints': 'points', 'categorie': 'category',
    },
    'client_ca': {
        'cartefidelite': 'card', 'totalca': 'amount', 'codemag': 'shop', 'date': 'date', 'client': 'fm_id',
        'nom': 'lastname', 'prenom': 'firstname',
    },
    'sales_detail': {
        'codemag': 'shop', 'vente': 'ticket', 'date': 'date', 'utilisateur': 'cashier', 'vendeur': 'seller',
        'client': 'fm_id', 'cartefidelite': 'card', 'nom': 'lastname', 'prenom': 'firstname', 'barcode': 'ref',
        'gencod': 'ean', 'reffournisseur': 'supplier_ref', 'designation': 'name', 'axe': 'axe', 'lignes': 'line',
        'quantite': 'qty', 'prix': 'price', 'remise': 'discount', 'total': 'total', 'motif': 'reason',
        'prixnet': 'net_price',
        'commentaire': 'comment',
    },
    'stores': {'magasin': 'name', 'nommagasin': 'name', 'code': 'code', 'codemagasin': 'code'},
    'sellers': {'codemagasin': 'shop', 'codemag': 'shop', 'magasin': 'shop', 'vendeurs': 'seller',
                'vendeur': 'seller', 'codelogistique': 'logistic_code'},
}
SIGNATURE = {
    'article_sage': {'article', 'categori', 'codeabarre'},
    'client_detail': {'client', 'cartefidelite', 'pubparemail'},
    'article': {'refarticle', 'eanbar'},
    'client_new': {'cardnumbre'},
    'client_old': {'numcliente'},
    'client_codes': {'codeclient', 'cartefidelite'},
    'client_ca': {'cartefidelite', 'totalca'},
    'sales_detail': {'vente', 'utilisateur', 'vendeur'},
    'stores': {'magasin', 'code'},
    'sellers': {'vendeurs'},
}
# civilities written in many ways in the files → the standard list (gender deduced when the file has none)
TITLE_ALIASES = [
    ('base.res_partner_title_madam', ('madame', 'mme', 'mrs', 'madam', 'mde', 'femme'), 'f'),
    ('base.res_partner_title_miss', ('mademoiselle', 'mlle', 'melle', 'mle', 'miss'), 'f'),
    ('base.res_partner_title_mister', ('monsieur', 'm', 'mr', 'mister', 'mons', 'homme'), 'h'),
    ('base.res_partner_title_doctor', ('docteur', 'dr', 'doctor'), False),
    ('base.res_partner_title_prof', ('professeur', 'pr', 'prof', 'professor'), False),
]
LAST_WINS = {'article_sage': {'brand'}}    # two columns with the same header: the last one holds the name
DEDUCED_REASON = 'Remise (déduite)'
FREE_REASON = 'Gratuité (déduite)'           # articles whose code starts with Y are free gifts
STAMP_REFS = {'DTT', 'TIMBRE'}      # « droit de timbre sur ticket »: not an article
MAX_NOTES = 100000                   # lines of the detailed report (Excel)


class StoreIndex(dict):
    """fastmag shop code → store. Unknown « YR_AZUR » codes are matched on the store name (AZUR, YOUG → YOUGOSLAVIE)
    when only one store fits, and the code is then saved on the store (« Code magasin fastmag »)."""

    def __init__(self, stores):
        super().__init__()
        self.records = stores
        for store in stores:
            for code in filter(None, [store.code] + (store.yr_fastmag_code or '').split(',')):
                self[code.strip().upper()] = store

    def _resolve(self, code):
        code = (code or '').strip().upper()
        if not code:
            return None
        if dict.__contains__(self, code):
            return dict.__getitem__(self, code)
        key = norm(re.sub(r'^YR[_\s-]*', '', code, flags=re.I))
        found = self.records.browse()
        if key:
            found = self.records.filtered(lambda s: norm(s.code) == key) or \
                self.records.filtered(lambda s: norm(s.name) == key) or \
                self.records.filtered(lambda s: norm(s.name).startswith(key))
            if len(found) != 1:
                found = self._closest(key.replace('center', 'centre'))      # SCENTER: English spelling
        if len(found) != 1:
            self[code] = None
            return None
        found.sudo().yr_fastmag_code = ','.join(filter(None, [found.yr_fastmag_code, code]))
        self[code] = found
        return found

    def _closest(self, key):
        """Typing errors and abbreviations: MAYHSOUNA → MAHSOUNA, HB → Habib Bourguiba, SCENTER → Sousse Centre.
        A store already known under another YR_ code (Sfax Centre = YR_SFAX_C) is not a candidate, and only a
        clear winner is kept."""
        scored = []
        for store in self.records.filtered(lambda s: s.store_type != 'website' or 'shop' in key or 'web' in key):
            own = [norm(re.sub(r'^YR[_\s-]*', '', alias.strip(), flags=re.I))
                   for alias in (store.yr_fastmag_code or '').split(',') if alias.strip().upper().startswith('YR')]
            if own:
                # already known under its own fastmag code: only a typing error of that code (MAYHSOUNA)
                scored.append((max(difflib.SequenceMatcher(None, key, form).ratio() for form in own), store))
                continue
            words = [norm(w) for w in re.split(r'[\s_\-/]+', store.name or '') if norm(w)]
            forms = {norm(store.name), ''.join(w[:1] for w in words)}        # full name, initials (HB)
            for i in range(1, len(words)):          # first words abbreviated: « S CENTRE », « SF CENTRE »
                for size in (1, 2, 3):
                    forms.add(''.join(w[:size] for w in words[:i]) + ''.join(words[i:]))
            scored.append((max(difflib.SequenceMatcher(None, key, form).ratio() for form in forms), store))
        scored.sort(key=lambda sc: -sc[0])
        if scored and scored[0][0] >= 0.8 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.08):
            return scored[0][1]
        return self.records.browse()

    def get(self, code, default=None):
        return self._resolve(code) or default

    def __contains__(self, code):
        return bool(self._resolve(code))

    def __getitem__(self, code):
        store = self._resolve(code)
        if not store:
            raise KeyError(code)
        return store


def norm(text):
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode()
    return ''.join(c for c in text.lower() if c.isalnum())


def text(value):
    """Cell → clean string (1006.0 → '1006', 104065941 → '104065941')."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    return ' '.join(str(value).split())


def number(value):
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(' ', '').replace(',', '.'))
    except ValueError:
        return None


def parse_date(value):
    """datetime / 'jj/mm/aaaa' / 'aaaa-mm-jj' → date, None if empty or not a real date (00/01/0000)."""
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def parse_birthday(value, unknown_years=()):
    """→ (day, month, year or None). '20/11/0000' = day and month only; '00/01/0000' = nothing."""
    if value in (None, ''):
        return None
    if isinstance(value, (datetime, date)):
        day, month, year = value.day, value.month, value.year
    else:
        match = re.match(r'^\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{1,4})', str(value))
        if not match:
            return None
        day, month, year = (int(g) for g in match.groups())
    try:
        date(2000, month, day)          # leap year: 29/02 accepted
    except ValueError:
        return None
    if year in unknown_years or not 1900 <= year <= date.today().year - 5:
        year = None
    return day, month, year


def card_number(value):
    """Loyalty card number, '' when it is not a real card (empty, 0, 000000, « numCliente »…)."""
    value = text(value).replace(' ', '')
    digits = re.sub(r'\D', '', value)
    if len(digits) < 6 or not digits.strip('0'):
        return ''
    return value


def client_code(value):
    """fastmag customer code: 'Y88074', or the number alone (88074, as in the CA files) → 'Y88074'."""
    value = text(value).replace(' ', '').upper()
    if value.isdigit():
        return 'Y%s' % value if value.strip('0') else ''
    return value if re.match(r'^Y\d+$', value) else ''


def is_coupon(item):
    """A coupon line: the word « coupon » in its label (or reference / motif). Codes (Y13051967…) are not a
    criterion: some free products carry such codes."""
    return any('coupon' in norm(item.get(key)) for key in ('name', 'ref', 'ean', 'reason'))


def is_xls(content):
    return content[:4] == b'\xd0\xcf\x11\xe0'


def _xlrd():
    try:
        import xlrd
    except ImportError:
        raise UserError(_('Fichier Excel ancien (.xls) : enregistrez-le en .xlsx (Fichier › Enregistrer sous) '
                          'puis réimportez-le.'))
    return xlrd


def xls_sheet_names(content):
    return _xlrd().open_workbook(file_contents=content, on_demand=True).sheet_names()


def xls_count(content, sheet=None):
    book = _xlrd().open_workbook(file_contents=content, on_demand=True)
    return (book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)).nrows


def xls_rows(content, sheet=None):
    """Rows of an old Excel file (.xls): dates become datetimes, whole numbers integers."""
    xlrd = _xlrd()
    book = xlrd.open_workbook(file_contents=content, on_demand=True)
    worksheet = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
    for r in range(worksheet.nrows):
        row = []
        for cell in worksheet.row(r):
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                except Exception:  # noqa: BLE001
                    value = None
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                value = None
            elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                value = int(value)
            row.append(value)
        yield row


def phone(value):
    digits = re.sub(r'\D', '', text(value))
    return digits or ''


class YrMigrationImport(models.TransientModel):
    _name = 'yr.migration.import'
    _description = 'Import fastmag'

    file = fields.Binary('Fichier Excel ou CSV')
    file_ids = fields.Many2many(
        'ir.attachment', 'yr_migration_import_attachment_rel', 'wizard_id', 'attachment_id', 'Plusieurs fichiers',
        help='Les fichiers sont importés l\'un après l\'autre : d\'abord par type (magasins, articles, clients, '
             'ventes), puis par numéro de partie du nom (ventes_partie_2 avant ventes_partie_10).')
    skip_done = fields.Boolean(
        'Ignorer les fichiers déjà importés', default=True,
        help='Un fichier déjà importé en entier (même contenu) est passé. Décochez pour le réimporter : les lignes '
             'sont mises à jour, sans doublon.')
    filename = fields.Char()
    file_type = fields.Selection(FILE_TYPES, 'Type de fichier', compute='_compute_file_type', store=True,
                                 readonly=False, help='Reconnu automatiquement d\'après les en-têtes.')
    # articles
    create_categories = fields.Boolean('Créer les catégories PdV manquantes (axe › sous-axe)', default=True)
    import_no_card = fields.Boolean(
        'Importer aussi les clients sans n° de carte', default=True,
        help='Un client sans vrai numéro de carte (vide, 0, 000000…) est importé avec son code client (Y…) : ses ventes '
             'du détail commercial lui sont rattachées par ce code. Il n\'a pas de numéro de carte fidélité. Une carte '
             'partagée par plusieurs clients reste sur la première fiche : les autres sont importés sans carte.')
    deduce_motif = fields.Boolean(
        'Déduire le motif des lignes remisées', default=True,
        help='L\'historique ne dit pas comment une promotion a été faite. Une ligne remisée (au moins 1 %) sans motif '
             'reçoit le motif « Remise (déduite) », pour qu\'elle compte dans les statistiques de promotion.')
    archive_inactive = fields.Boolean('Archiver les articles au statut autre que « Actif »', default=False)
    update_prices = fields.Boolean('Mettre à jour les prix des articles existants', default=True)
    # customers
    year_2000_unknown = fields.Boolean(
        'Année 2000 = année inconnue (fichier ancien système)', default=True,
        help='Dans l\'ancien système, beaucoup de dates de naissance ont l\'année 2000 par défaut : seuls le '
             'jour et le mois sont alors repris.')
    update_existing = fields.Boolean('Compléter les clients existants', default=True,
                                     help='Les champs vides de la fiche sont complétés ; rien n\'est effacé.')
    # sales detail
    create_sellers = fields.Boolean(
        'Créer les vendeurs manquants (Employés)', default=False,
        help='Les noms « Vendeur » et « Caissier » de fastmag sont toujours gardés sur chaque ligne de vente. '
             'Coché : chaque vendeur inconnu devient aussi un employé. Décoché : seuls les noms déjà connus '
             '(employé ou « Nom(s) fastmag ») sont rattachés à un employé.')
    create_pos_orders = fields.Boolean(
        'Créer les commandes du point de vente', default=True,
        help='Chaque ticket devient une commande de Point de Vente › Commandes (session « historique fastmag » '
             'fermée de la caisse du magasin, sans paiement, écriture comptable ni mouvement de stock). '
             'Créées en arrière-plan après l\'import.')
    update_purchases = fields.Boolean(
        'Alimenter l\'historique d\'achats des clients', default=True,
        help='Total par carte, jour et magasin, comme le fichier « CA clients » (sans doublon avec lui) : '
             'statut client, CA, segments de Fidélité & CRM.')
    # stores
    store_companies = fields.Boolean(
        'Une société (filiale) par magasin', default=True,
        help='Chaque magasin devient une filiale de la société principale (ventes, stock, vendeurs, statistiques '
             'par magasin). Le site web est une filiale « Site web ».')
    store_pos = fields.Boolean(
        'Un point de vente (caisse) par magasin', default=True,
        help='Chaque magasin physique ouvert reçoit sa caisse (et un entrepôt dans sa filiale) : les ventes '
             'importées y sont rattachées.')
    state = fields.Selection([('draft', 'draft'), ('done', 'done')], default='draft')
    log_id = fields.Many2one('yr.migration.log', 'Import', readonly=True)
    previous_import = fields.Html('Déjà importé', compute='_compute_previous_import')
    result = fields.Html('Résultat', readonly=True)

    # ── history, progress, resume ────────────────────────────────────────
    def _content(self):
        """The file itself (the screen reads binary fields as their size only)."""
        return base64.b64decode(self.with_context(bin_size=False).file or b'')

    def _file_hash(self):
        return hashlib.sha1(self._content()).hexdigest()

    def _previous_log(self):
        return self.env['yr.migration.log'].search([('file_hash', '=', self._file_hash())], limit=1) \
            if self.file else self.env['yr.migration.log']

    @api.depends('file')
    def _compute_previous_import(self):
        for wizard in self:
            log = wizard._previous_log()
            if not log:
                wizard.previous_import = False
            elif log.state == 'done':
                wizard.previous_import = _(
                    '<p class="text-info">Ce fichier a déjà été importé le %(date)s (%(done)s lignes). Le réimporter '
                    'met à jour les données, sans doublon.</p>', date=fields.Datetime.to_string(log.date_end or
                                                                                               log.date_start),
                    done=log.done)
            else:
                wizard.previous_import = _(
                    '<p class="text-warning">Import de ce fichier interrompu le %(date)s à %(done)s / %(total)s lignes. '
                    '« Importer » reprend où il s\'est arrêté (les lignes déjà enregistrées ne sont pas refaites).</p>',
                    date=fields.Datetime.to_string(log.date_start), done=log.done, total=log.total)

    def _can_commit(self):
        # not inside the automated tests (one transaction, rolled back)
        return not odoo_module.current_test and not self.env.registry.in_test_mode()

    def _commit(self):
        """Saves what is imported so far: an interrupted import resumes from there."""
        if self._can_commit():
            self.env.cr.commit()

    def _progress(self, done, total, label, force=False):
        """Live progress (bar of the import screen), written apart from the import transaction."""
        log = self.log_id
        if not log:
            return
        now = time.monotonic()
        if not force and now - self.env.context.get('yr_progress_clock', [0])[0] < 1.0:
            return
        self.env.context.get('yr_progress_clock', [0])[0] = now
        if not self._can_commit():
            log.sudo().write({'done': done, 'total': total, 'label': label})
            return
        with self.env.registry.cursor() as cr:
            cr.execute('UPDATE yr_migration_log SET done = %s, total = %s, label = %s WHERE id = %s',
                       (done, total, label, log.id))

    def get_progress(self):
        """Polled by the import screen while the import runs."""
        log = self.sudo().log_id
        if not log:
            return {}
        return {'state': log.state, 'done': log.done, 'total': log.total, 'label': log.label or '',
                'progress': log.progress}

    # ── reading ──────────────────────────────────────────────────────────
    def _rows(self, raw=False):
        """Header (normalised, or as is with raw=True) + data rows, empty rows skipped."""
        content = self._content()
        name = (self.filename or '').lower()
        if is_xls(content):                       # old Excel file (.xls)
            sheet = self.env.context.get('yr_sheet')
            rows = xls_rows(content, sheet)
            total = xls_count(content, sheet)
        elif name.endswith(('.xlsx', '.xlsm')) or content[:2] == b'PK':
            import openpyxl
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = self.env.context.get('yr_sheet')
            worksheet = workbook[sheet] if sheet else workbook.active
            total = worksheet.max_row or 0
            rows = worksheet.iter_rows(values_only=True)
        else:
            decoded = ''
            for encoding in ('utf-8-sig', 'cp1252', 'latin-1'):
                try:
                    decoded = content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            delimiter = max([';', ',', '\t', '|'], key=decoded[:4096].count)
            total = decoded.count('\n')
            rows = csv.reader(io.StringIO(decoded), delimiter=delimiter)
        header = None
        reporting = not raw and self.env.context.get('yr_report_reading')
        for count, row in enumerate(rows, start=1):
            if reporting and count % 500 == 0:
                self._progress(count, total, _('Lecture du fichier'))
            if not any(c not in (None, '') and str(c).strip() for c in row):
                continue
            if header is None:
                header = [norm(c) for c in row]
                yield list(row) if raw else header
                continue
            yield list(row)
        if header is None:
            raise UserError(_('Le fichier est vide.'))

    def _sheet_names(self):
        """Sheets of an Excel workbook (empty for a CSV file)."""
        content = self._content()
        if is_xls(content):
            return xls_sheet_names(content)
        if not ((self.filename or '').lower().endswith(('.xlsx', '.xlsm')) or content[:2] == b'PK'):
            return []
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(content), read_only=True).sheetnames

    def _detect_kind(self):
        """File type of the current sheet (context ``yr_sheet``) from its header."""
        try:
            header = set(next(self._rows()))
        except Exception:  # noqa: BLE001 - unreadable / empty sheet
            return False
        if 'nummagasin' in header and 'magasin' in header:
            return 'stores'                              # magasin | code fastmag (YR_CARR) | numéro (2)
        return next((kind for kind, keys in SIGNATURE.items() if keys <= header), False) or \
            ('stores' if self._is_store_list() else False)

    def _sheet_kinds(self):
        """[(sheet, file type)] of the recognised sheets of a workbook, in import order."""
        kinds = []
        for sheet in self._sheet_names():
            kind = self.with_context(yr_sheet=sheet)._detect_kind()
            if kind:
                kinds.append((sheet, kind))
        return sorted(kinds, key=lambda sk: SHEET_ORDER.index(sk[1]))

    @api.depends('file', 'filename')
    def _compute_file_type(self):
        for wizard in self:
            wizard.file_type = False
            if not wizard.file:
                continue
            try:
                sheets = wizard._sheet_kinds()
            except Exception:  # noqa: BLE001 - unreadable file: reported when importing
                sheets = []
            if len(sheets) > 1:
                wizard.file_type = 'workbook'
            elif sheets:
                wizard.file_type = sheets[0][1]
            else:
                wizard.file_type = wizard._detect_kind()

    def _is_store_list(self):
        """A list without header row: store name, numeric code (« CARREFOUR | 2 »)."""
        try:
            first = [c for c in next(self._rows(raw=True)) if c not in (None, '')]
        except Exception:  # noqa: BLE001
            return False
        return len(first) == 2 and number(first[1]) is not None and number(first[0]) is None

    def _records(self, kind):
        rows = self._rows()
        header = next(rows)
        mapping = {}
        for idx, col in enumerate(header):
            key = COLUMNS[kind].get(col)
            if key and (key not in mapping or key in LAST_WINS.get(kind, ())):
                mapping[key] = idx
        missing = [col for col in SIGNATURE[kind] if col not in header]
        if missing:
            raise UserError(_('Ce fichier n\'est pas un fichier « %s » : colonne(s) %s introuvable(s).',
                              dict(FILE_TYPES)[kind], ', '.join(missing)))
        for line_no, row in enumerate(rows, start=2):
            item = {key: (row[idx] if idx < len(row) else None) for key, idx in mapping.items()}
            item['_line'] = line_no
            yield item

    # ── entry point ──────────────────────────────────────────────────────
    def action_import(self):
        self.ensure_one()
        if not self.file and self.file_ids:
            return self.action_import_queue()
        if not self.file:
            raise UserError(_('Choisissez un fichier.'))
        if not self.file_type:
            raise UserError(_('Type de fichier non reconnu : choisissez-le dans la liste.'))
        previous = self._previous_log()
        resume = bool(previous) and previous.state != 'done'
        notes = []                           # detailed report: (part, line, code, result, reason)
        log = previous if resume else self.env['yr.migration.log'].create({
            'name': self.filename or _('Fichier'), 'file_hash': self._file_hash(), 'file_type': self.file_type})
        log.write({'state': 'running', 'label': _('Démarrage'), 'user_id': self.env.uid})
        self.log_id = log
        item = self.env['yr.migration.queue'].browse(self.env.context.get('yr_queue_item') or [])
        if item:
            item.log_id = log                # the queue screen follows the file being imported
        self._commit()                       # the screen can follow the progress from now on
        env = self.with_context(tracking_disable=True, mail_create_nolog=True, mail_notrack=True,
                                yr_migration_force=True, yr_resume=resume, yr_progress_clock=[0], yr_cache={},
                                yr_report_reading=True, yr_notes=notes)
        if self.file_type == 'workbook':
            parts = [(sheet, kind) for sheet, kind in self._sheet_kinds()]
        else:
            sheets = [sheet for sheet, kind in self._sheet_kinds() if kind == self.file_type]
            parts = [(sheets[0] if sheets else None, self.file_type)]
        result = ''
        summary = []
        try:
            for sheet, kind in parts:
                stats = defaultdict(int)
                errors = []
                if resume:
                    stats['resumed'] = 1
                title = self._part_title(kind, sheet if self.file_type == 'workbook' else None)
                getattr(env.with_context(yr_sheet=sheet, yr_part=title), '_import_%s' % kind)(stats, errors)
                result += self._render(stats, errors, kind, sheet if self.file_type == 'workbook' else None)
                summary += [(title, label, stats[key]) for key, label in self._stat_labels().items()
                            if stats.get(key)]
                for line, code, message in errors:
                    if len(notes) < MAX_NOTES:
                        notes.append((title, line, code or '', _('Erreur'), message))
        except Exception:
            if self._can_commit():
                self.env.cr.rollback()
                with self.env.registry.cursor() as cr:
                    cr.execute("UPDATE yr_migration_log SET state = 'failed' WHERE id = %s", (log.id,))
            raise
        # the progress was written beside this transaction: a fresh one sees it (no concurrent-update error)
        self._commit()
        log.invalidate_recordset()
        report_name = 'rapport_import_%s.xlsx' % re.sub(r'\W+', '_', (self.filename or 'fastmag').rsplit('.', 1)[0])
        result += '<p><a href="/web/content/yr.migration.log/%s/report_file/%s?download=true" class="btn btn-link ' \
                  'p-0"><i class="fa fa-download"/> %s</a></p>' % (log.id, report_name, html_escape(
                      _('Télécharger le rapport détaillé (Excel) : résumé et lignes non importées')))
        log.write({'state': 'done', 'result': result, 'date_end': fields.Datetime.now(), 'label': _('Terminé'),
                   'done': max(log.done, log.total), 'total': max(log.done, log.total), 'report_filename': report_name,
                   'report_file': base64.b64encode(self._report_xlsx(summary, notes))})
        self.write({'state': 'done', 'result': result})
        return {'type': 'ir.actions.act_window', 'name': _('Reprise fastmag'), 'res_model': self._name,
                'res_id': self.id, 'view_mode': 'form', 'target': 'new'}

    # ── several files, one after the other ───────────────────────────────
    OPTIONS = ('create_categories', 'update_prices', 'archive_inactive', 'create_pos_orders', 'deduce_motif',
               'import_no_card', 'create_sellers', 'update_purchases', 'store_companies', 'store_pos',
               'year_2000_unknown', 'update_existing')

    @staticmethod
    def _part_number(filename):
        """Part number of a file name: the number after « part / partie / p », else the last number."""
        stem = re.sub(r'\.\w+$', '', filename or '')
        found = re.findall(r'(?:parti?e?|pt|p)[\s._-]*(\d+)', stem, re.I) or re.findall(r'(\d+)', stem)
        return int(found[-1]) if found else 0

    def action_import_queue(self):
        """Puts the chosen files in the import queue: they are imported one after the other in the background, by
        file type (the order the data needs) then by part number, without anyone in front of the screen."""
        self.ensure_one()
        options = json.dumps({name: self[name] for name in self.OPTIONS})
        Queue = self.env['yr.migration.queue']
        for attachment in self.file_ids:
            kind = self.create({'file': attachment.datas, 'filename': attachment.name}).file_type
            order = SHEET_ORDER.index(kind) if kind in SHEET_ORDER else (-1 if kind == 'workbook' else len(SHEET_ORDER))
            Queue.create({'name': attachment.name, 'attachment_id': attachment.id, 'options': options,
                          'skip_done': self.skip_done, 'file_type': kind or False, 'type_order': order,
                          'part': self._part_number(attachment.name)})
        self.file_ids = [(5, 0, 0)]
        self._commit()
        Queue._trigger_cron()
        action = self.env['ir.actions.act_window']._for_xml_id('yr_data_migration.action_yr_migration_queue')
        action['target'] = 'current'
        return action

    def _note(self, line, code, result, reason):
        """One line of the detailed report (Excel)."""
        notes = self.env.context.get('yr_notes')
        if notes is not None and len(notes) < MAX_NOTES:
            notes.append((self.env.context.get('yr_part') or '', line, code or '', result, reason))

    def _part_title(self, kind, sheet=None):
        title = dict(FILE_TYPES)[kind or self.file_type]
        return '%s — %s' % (sheet, title) if sheet else title

    def _report_xlsx(self, summary, notes):
        """Excel report of the import: summary per file / sheet, then every line ignored or in error."""
        import openpyxl
        workbook = openpyxl.Workbook(write_only=True)
        sheet = workbook.create_sheet(_('Résumé'))
        sheet.append([_('Fichier'), self.filename or ''])
        sheet.append([_('Date'), fields.Datetime.to_string(fields.Datetime.now())])
        sheet.append([_('Par'), self.env.user.name])
        sheet.append([])
        sheet.append([_('Partie'), _('Indicateur'), _('Nombre')])
        for title, label, value in summary:
            sheet.append([title, label, ', '.join(sorted(value)) if isinstance(value, set) else value])
        detail = workbook.create_sheet(_('Détail'))
        detail.append([_('Partie'), _('Ligne'), _('Code'), _('Résultat'), _('Motif')])
        for note in notes:
            detail.append(list(note))
        if len(notes) >= MAX_NOTES:
            detail.append(['', '', '', '', _('… liste limitée à %s lignes', MAX_NOTES)])
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def _stat_labels(self):
        return {
            'rows': _('Lignes lues'), 'created': _('Créés'), 'updated': _('Mis à jour'), 'skipped': _('Ignorés'),
            'categories': _('Catégories PdV créées'), 'suppliers': _('Fournisseurs créés'),
            'archived': _('Articles archivés'), 'unknown_cards': _('Cartes sans fiche client'),
            'unknown_stores': _('Lignes avec magasin inconnu'), 'customers': _('Clients avec historique'),
            'tickets': _('Tickets'), 'purchases': _('Achats clients (historique CRM)'),
            'employees': _('Vendeurs créés (Employés)'), 'unknown_products': _('Lignes sans article connu'),
            'sellers_updated': _('Vendeurs mis à jour'), 'closed': _('Magasins fermés (archivés)'),
            'companies': _('Sociétés (filiales) créées'), 'pos': _('Points de vente (caisses) créés'),
            'already': _('Lignes déjà importées (reprise : non refaites)'),
            'added_to_known': _('Lignes manquantes ajoutées à des ventes déjà importées'),
            'coupons': _('Lignes de coupon / bon d\'achat (non importées)'),
            'no_card': _('Clients sans vrai n° de carte fidélité (non importés)'),
            'imported_no_card': _('Clients importés sans n° de carte (rattachés à leurs ventes par le code client)'),
            'shared_cards': _('Cartes partagées : clients importés sans carte (la carte reste sur la première fiche)'),
            'codes': _('Codes clients fastmag'), 'codes_no_card': _('Codes sans vraie carte (ventes anonymes)'),
            'codes_linked': _('Codes rattachés à un client'),
            'relinked': _('Lignes de ventes déjà importées rattachées à leur client'),
            'via_code': _('Lignes rattachées au client par le code Y'),
            'unmapped_code': _('Lignes avec code client sans carte connue (sans client)'),
            'anonymous': _('Lignes sans client (anonymes)'),
            'sellers_seen': _('Vendeurs (noms fastmag)'), 'cashiers_seen': _('Caissiers (noms fastmag)'),
            'pos_pending': _('Tickets envoyés vers Point de Vente › Commandes (création en arrière-plan)'),
        }

    def _render(self, stats, errors, kind=None, sheet=None):
        labels = self._stat_labels()
        title = self._part_title(kind, sheet)
        html = '<p><b>%s</b></p><ul>' % html_escape(title)
        for key, label in labels.items():
            if stats.get(key):
                html += '<li>%s : <b>%s</b></li>' % (html_escape(label), stats[key])
        html += '</ul>'
        if stats.get('stores_created'):
            html += '<p class="text-info">%s %s</p>' % (
                html_escape(_('Magasins créés (inconnus jusqu\'ici, à compléter dans Fidélité & CRM › Configuration › '
                              'Magasins) :')), html_escape(', '.join(sorted(stats['stores_created']))))
        if stats.get('stores_missing'):
            html += '<p class="text-warning">%s %s</p>' % (
                html_escape(_('Codes magasin à renseigner dans Fidélité & CRM › Configuration › Magasins '
                              '(champ « Code magasin fastmag ») puis réimporter :')),
                html_escape(', '.join(sorted(stats['stores_missing']))))
        if errors:
            html += '<p class="text-danger"><b>%s</b></p><ul>' % html_escape(_('%s ligne(s) en erreur', len(errors)))
            for line, code, message in errors[:200]:
                html += '<li>%s %s — %s : %s</li>' % (html_escape(_('Ligne')), line, html_escape(code or ''),
                                                       html_escape(message))
            html += '</ul>'
        return html

    # ── articles ─────────────────────────────────────────────────────────
    def _pos_category(self, name, parent, cache, stats):
        name = text(name)
        if not name:
            return self.env['pos.category']
        key = (name.upper(), parent.id)
        if key not in cache:
            Category = self.env['pos.category']
            categ = Category.search([('name', '=ilike', name), ('parent_id', '=', parent.id or False)], limit=1)
            if not categ and self.create_categories:
                categ = Category.create({'name': name, 'parent_id': parent.id or False,
                                         'is_crm_family': not parent, 'import_code': name.upper()})
                stats['categories'] += 1
            cache[key] = categ
        return cache[key]

    def _supplier(self, code, name, cache, stats):
        code, name = text(code), text(name)
        if not code and not name:
            return self.env['res.partner']
        key = code or name
        if key not in cache:
            Partner = self.env['res.partner'].with_context(crm_no_card=True)
            supplier = Partner.search([('ref', '=', code), ('is_company', '=', True)], limit=1) if code else Partner
            if not supplier and name:
                supplier = Partner.search([('name', '=ilike', name), ('is_company', '=', True)], limit=1)
            if not supplier:
                supplier = Partner.create({'name': name or code, 'ref': code or False, 'is_company': True,
                                           'crm_is_customer': False, 'supplier_rank': 1})
                stats['suppliers'] += 1
            cache[key] = supplier
        return cache[key]

    def _import_article_sage(self, stats, errors):
        self._import_article(stats, errors, 'article_sage')

    def _import_article(self, stats, errors, kind='article'):
        items = []
        for item in self._records(kind):
            if not (text(item.get('code')) or text(item.get('barcode'))):
                continue
            if kind == 'article_sage':
                # Etat 1 = actif; Nature « 0 » = empty
                item['status'] = 'Actif' if text(item.get('status')) in ('1', '') else 'Inactif'
                if text(item.get('nature')) == '0':
                    item['nature'] = ''
            items.append(item)
        stats['rows'] = len(items)
        Template = self.env['product.template'].with_context(active_test=False)
        codes = [text(i.get('code')) for i in items if text(i.get('code'))]
        barcodes = [text(i.get('barcode')) for i in items if text(i.get('barcode'))]
        by_code, by_barcode = {}, {}
        for chunk in range(0, max(len(codes), len(barcodes)), 5000):
            for tmpl in Template.search(['|', ('default_code', 'in', codes[chunk:chunk + 5000]),
                                         ('barcode', 'in', barcodes[chunk:chunk + 5000])]):
                by_code.setdefault(tmpl.default_code, tmpl)
                by_barcode.setdefault(tmpl.barcode, tmpl)
        categ_cache, supplier_cache = {}, {}
        unit = self.env.ref('uom.product_uom_unit')
        seen_barcodes = set()
        to_create = []
        for position, item in enumerate(items, start=1):
            if position % 200 == 0:
                self._progress(position, len(items), _('Enregistrement des articles'))
            code, barcode = text(item.get('code')), text(item.get('barcode'))
            owner = by_barcode.get(barcode) if barcode else None
            # a line with a reference matches on it: two fastmag references may share a barcode
            if code:
                tmpl = by_code.get(code)
                if not tmpl and owner and owner is not True and not owner.default_code:
                    tmpl = owner                       # article created before without reference
            else:
                tmpl = owner
            if barcode and (barcode in seen_barcodes or (by_barcode.get(barcode) and by_barcode[barcode] != tmpl)):
                errors.append((item['_line'], code, _('code-barres %s déjà utilisé par un autre article : '
                                                     'article repris sans code-barres', barcode)))
                barcode = ''
            if barcode:
                seen_barcodes.add(barcode)
            axe = self._pos_category(item.get('axe'), self.env['pos.category'], categ_cache, stats)
            sub_axe = self._pos_category(item.get('sub_axe'), axe, categ_cache, stats) if axe else axe
            price = number(item.get('price'))
            status = text(item.get('status'))
            vals = {
                'yr_site': text(item.get('site')) or False,
                'yr_category_code': text(item.get('category')) or False,
                'yr_brand': text(item.get('brand')) or False,
                'yr_line': text(item.get('line')) or False,
                'yr_axe_id': axe.id or False,
                'yr_sub_axe_id': sub_axe.id or False,
                'yr_nature': text(item.get('nature')) or False,
                'yr_packaging': text(item.get('packaging')) or False,
                'yr_gender': text(item.get('gender')) or False,
                'yr_status': status or False,
                'yr_create_date_legacy': parse_date(item.get('create_date')) or False,
                'yr_supplier_code': text(item.get('supplier_code')) or False,
            }
            if sub_axe or axe:
                vals['pos_categ_ids'] = [(6, 0, (sub_axe or axe).ids)]
            if self.archive_inactive and status:
                vals['active'] = status.lower().startswith('actif')
                if not vals['active']:
                    stats['archived'] += 1
            supplier = self._supplier(item.get('supplier_code'), item.get('supplier_name'), supplier_cache, stats)
            if tmpl is True:          # same reference twice in the file: first line wins
                stats['skipped'] += 1
                continue
            if tmpl:
                if code and not tmpl.default_code:
                    vals['default_code'] = code
                if barcode and not tmpl.barcode:
                    vals['barcode'] = barcode
                if price is not None and self.update_prices:
                    vals['list_price'] = price
                if text(item.get('name')):
                    vals['name'] = text(item.get('name'))
                tmpl.write(vals)
                if supplier and supplier not in tmpl.seller_ids.partner_id:
                    tmpl.seller_ids = [(0, 0, {'partner_id': supplier.id})]
                stats['updated'] += 1
            else:
                vals.update({
                    'name': text(item.get('name')) or code or barcode, 'default_code': code or False,
                    'barcode': barcode or False, 'list_price': price or 0.0, 'type': 'consu',
                    'is_storable': True, 'available_in_pos': True, 'sale_ok': True,
                    'uom_id': unit.id, 'uom_po_id': unit.id,
                })
                if supplier:
                    vals['seller_ids'] = [(0, 0, {'partner_id': supplier.id})]
                to_create.append(vals)
                if code:
                    by_code[code] = True       # same code twice in the file: created once
            if len(to_create) >= BATCH:
                stats['created'] += len(Template.create(to_create))
                to_create = []
        if to_create:
            stats['created'] += len(Template.create(to_create))

    # ── customers ────────────────────────────────────────────────────────
    def _stores(self):
        # closed (archived) stores too: their old sales and sellers are still imported
        # every company (the stores are in their subsidiaries, whatever company is selected in the session)
        return StoreIndex(self.env['loyalty.store'].sudo().with_context(active_test=False).search([]))

    def _existing_customers(self, cards):
        Partner = self.env['res.partner'].with_context(active_test=False)
        found = {}
        for chunk in range(0, len(cards), 5000):
            part = cards[chunk:chunk + 5000]
            for partner in Partner.search(['|', ('customer_code', 'in', part), ('yr_legacy_code', 'in', part)]):
                found.setdefault(partner.customer_code, partner)
                found.setdefault(partner.yr_legacy_code, partner)
        return found

    def _customer_vals(self, item, kind, stores, stats):
        unknown_years = (2000,) if kind == 'client_old' and self.year_2000_unknown else ()
        name = ' '.join(filter(None, [text(item.get('lastname')), text(item.get('firstname'))]))
        # Rue → street, Rue 2 (and 3) → street2
        street2 = ' '.join(filter(None, [text(item.get('street2')), text(item.get('street3'))]))
        vals = {
            'name': name, 'street': text(item.get('street')), 'street2': street2,
            'city': text(item.get('city')), 'zip': text(item.get('zip')),
            'country_id': self._tunisia_id(),
            'email': text(item.get('email')).lower(), 'yr_seller_code': text(item.get('seller')),
            'yr_seller_id': self._seller_id(text(item.get('seller'))),
            'yr_legacy_old_code': text(item.get('old_code')), 'yr_legacy_segment': text(item.get('segment')),
            'yr_entry_date': parse_date(item.get('entry_date')),
            'yr_first_purchase_legacy': parse_date(item.get('first_purchase')),
        }
        tel, mobile = phone(item.get('phone')), phone(item.get('mobile'))
        if kind == 'client_old' and tel and tel[:1] in '2459' and len(tel) == 8:
            mobile, tel = tel, ''            # Tunisian mobile numbers start with 2, 4, 5 or 9
        vals.update(phone=tel, mobile=mobile)
        gender = text(item.get('gender')).lower()
        vals['crm_gender'] = 'f' if gender[:1] == 'f' else 'h' if gender[:1] in ('h', 'm') else False
        # civilité / ancienne civilité: from the list (Madame, Mlle, M.…), a new value is added to the list
        title = self._title_id(item.get('title'))
        vals['title'] = title
        if kind == 'client_detail':
            vals['country_id'] = self._country_id(item.get('country'))
            for key in ('optin_email', 'optin_sms'):
                vals[key] = bool(number(item.get(key)))
        vals['crm_old_title_id'] = self._title_id(item.get('old_title'))
        if title and not vals['crm_gender']:
            vals['crm_gender'] = self._cache().get('title_gender', {}).get(title, False)
        if kind == 'client_old' and item.get('npai') not in (None, ''):
            vals['yr_npai'] = bool(number(item.get('npai')))
        birthday = parse_birthday(item.get('birthday'), unknown_years)
        if kind == 'client_detail':             # day, month and year in three columns (0 = unknown)
            day, month, year = (int(number(item.get(k)) or 0) for k in ('birth_d', 'birth_m', 'birth_y'))
            birthday = parse_birthday('%02d/%02d/%04d' % (day, month, year)) if day and month else None
        if birthday:
            day, month, year = birthday
            if year:
                try:
                    date(year, month, day)
                except ValueError:                # 29/02 of a year that was not a leap year: the year is dropped
                    year = None
            vals.update(birth_day=str(day), birth_month=str(month), birth_year=str(year) if year else False)
        amount = number(item.get('amount'))
        if amount:
            vals.update(yr_legacy_amount=amount, yr_migration_date=fields.Date.context_today(self))
        shop = text(item.get('shop')).upper()
        if shop and shop != 'CENTRAL':          # CENTRAL = the head office, not a store
            if shop in stores:
                vals['crm_store_id'] = stores[shop].id
            else:
                stats['stores_missing'] = stats.get('stores_missing') or set()
                stats['stores_missing'].add(shop)
                stats['unknown_stores'] += 1
        return {k: v for k, v in vals.items() if v not in (None, '', False) or k == 'birth_year'}

    def _title_id(self, value):
        """Civility of the file (Madame, MME, Mlle, M., Monsieur, Dr…) → res.partner.title, created if new."""
        key = norm(value)
        if not key or re.search(r'\d', text(value)) or len(key) > 15:
            return False                        # empty, or an address typed in the civility column
        cache = self._cache()
        if 'titles' not in cache:
            Title = self.env['res.partner.title'].sudo()
            titles, genders = {}, {}
            for xmlid, aliases, gender in TITLE_ALIASES:
                title = self.env.ref(xmlid, raise_if_not_found=False)
                if title:
                    titles.update(dict.fromkeys(aliases, title.id))
                    if gender:
                        genders[title.id] = gender
            for lang in [code for code, _name in self.env['res.lang'].get_installed()]:
                for title in Title.with_context(lang=lang).search([]):
                    for label in (title.name, title.shortcut):
                        if norm(label):
                            titles.setdefault(norm(label), title.id)
            cache['titles'], cache['title_gender'] = titles, genders
        if key not in cache['titles']:
            cache['titles'][key] = self.env['res.partner.title'].sudo().create({'name': text(value).title()}).id
        return cache['titles'][key]

    def _cache(self):
        """Per-import cache (dict given in the context by action_import)."""
        cache = self.env.context.get('yr_cache')
        return cache if cache is not None else {}

    def _tunisia_id(self):
        cache = self._cache()
        if 'tunisia' not in cache:
            tunisia = self.env.ref('base.tn', raise_if_not_found=False)
            cache['tunisia'] = tunisia.id if tunisia else False
        return cache['tunisia']

    def _country_id(self, value):
        """Country of the file (empty / Tunisie → Tunisia; another known country name is kept)."""
        name = text(value)
        if norm(name) in ('', 'tunisie', 'tunisia', 'tn', 'tun'):
            return self._tunisia_id()
        cache = self._cache()
        if name.lower() not in cache.setdefault('countries', {}):
            country = self.env['res.country'].search(['|', ('name', '=ilike', name), ('code', '=ilike', name)],
                                                     limit=1)
            cache['countries'][name.lower()] = country.id or self._tunisia_id()
        return cache['countries'][name.lower()]

    def _seller_id(self, code):
        """Seller of the file (code logistique 703, or fastmag name) → employee."""
        if not code:
            return False
        cache = self._cache()
        if 'sellers' not in cache:
            Employee = self.env['hr.employee'].sudo().with_context(active_test=False)
            sellers = {}
            for employee in Employee.search([('yr_logistic_code', '!=', False)]):
                sellers[employee.yr_logistic_code.strip().upper()] = employee.id
            for key, employee in self.env['yr.legacy.sale.line']._employee_index().items():
                sellers.setdefault(key, employee.id)
            cache['sellers'] = sellers
        return cache['sellers'].get(' '.join(code.upper().split()), False)

    def _import_customers(self, kind, stats, errors):
        items = []
        for item in self._records(kind):
            card = card_number(item.get('card'))
            code = client_code(item.get('fm_id'))
            if not card and norm(item.get('card')) == 'numcliente':
                continue                        # the sample's description row (« numCliente », « nom »…)
            name = ' '.join(filter(None, [text(item.get('lastname')), text(item.get('firstname'))]))
            if not card and not (self.import_no_card and code):
                # no real card and nothing else to recognise the customer by
                stats['no_card'] += 1
                self._note(item['_line'], text(item.get('card')), _('Non importé'),
                           _('pas de numéro de carte fidélité ni de code client (%s)', name or _('sans nom')))
                continue
            item['card'] = card
            item['code'] = code
            items.append(item)
        stats['rows'] = len(items) + stats['no_card']
        stores = self._stores()
        cards_in_file = list({i['card'] for i in items if i['card']})
        existing = self._existing_customers(cards_in_file)
        codes = self.env['yr.client.code'].sudo()._codes_for(set(cards_in_file))
        by_code = self._customers_by_code([i['code'] for i in items if i['code']])
        Partner = self.env['res.partner'].sudo().with_context(crm_keep_without_card=True)
        to_create, created_cards, created_codes = [], set(), set()
        for position, item in enumerate(items, start=1):
            if position % 200 == 0:           # also when the customers already exist (updates, skipped)
                self._progress(position, len(items), _('Enregistrement des clients'))
            card = item['card']
            vals = self._customer_vals(item, kind, stores, stats)
            code = item['code'] or codes.get(card)
            if code:
                vals['yr_client_code'] = code
            if not vals.get('birth_day'):
                vals.pop('birth_year', None)
            # the customer already known: by its card, else by its fastmag client code
            partner = existing.get(card) if card else None
            if not partner and code:
                partner = by_code.get(code)
            label = card or code
            if partner:
                if not self.update_existing:
                    stats['skipped'] += 1
                    continue
                # complete the empty fields only; the frozen migration amount is never replaced
                update = {k: v for k, v in vals.items() if not partner.sudo()[k] and k != 'birth_year'}
                # imported before with Rue 2 / Rue 3 glued to the street: split again
                glued = ' '.join(filter(None, [vals.get('street'), vals.get('street2')]))
                if vals.get('street2') and partner.street == glued and partner.street != vals.get('street'):
                    update.update(street=vals['street'], street2=vals['street2'])
                if 'birth_day' in update and vals.get('birth_year'):
                    update['birth_year'] = vals['birth_year']
                if 'birth_day' in update and 'birth_month' not in update:
                    update.pop('birth_day')
                if card and not partner.yr_legacy_code and existing.get(card) == partner:
                    update['yr_legacy_code'] = card
                if not partner.active:
                    update['active'] = True             # archived by « Supprimer les clients »: back
                if update:
                    try:
                        with self.env.cr.savepoint():
                            partner.sudo().write(update)
                        stats['updated'] += 1
                    except Exception as error:  # noqa: BLE001 - reported per line
                        errors.append((item['_line'], label, str(error)))
                else:
                    stats['skipped'] += 1
                continue
            if code and code in created_codes:
                errors.append((item['_line'], label, _('code client en double dans le fichier : ligne ignorée')))
                continue
            vals.update({'name': vals.get('name') or _('Client %s', label), 'crm_is_customer': True, 'customer_rank': 1})
            if card and card not in created_cards:
                vals.update(customer_code=card, barcode=card, yr_legacy_code=card)
                created_cards.add(card)
            elif card:
                # the card is already on another customer of the file: this one is imported without a card number
                if not code:
                    errors.append((item['_line'], label, _('carte déjà utilisée et pas de code client : ligne ignorée')))
                    continue
                stats['shared_cards'] += 1
                vals['comment'] = _('Carte %s partagée avec un autre client (gardée sur la première fiche).', card)
                self._note(item['_line'], card, _('Importé sans carte'),
                           _('carte partagée avec un autre client : importé sans n° de carte (%s)', code))
            else:
                stats['imported_no_card'] += 1
            if code:
                created_codes.add(code)
            to_create.append((item['_line'], vals))
            if len(to_create) >= BATCH:
                self._create_customers(Partner, to_create, stats, errors)
                to_create = []
                self._commit()                          # an interruption resumes from here
                self._progress(stats['created'] + stats['updated'] + stats['skipped'], len(items),
                               _('Enregistrement des clients'))
        if to_create:
            self._create_customers(Partner, to_create, stats, errors)
        cards = list({i['card'] for i in items if i['card']})
        codes_in_file = list({i['code'] for i in items if i['code']})
        self._progress(len(items), len(items), _('Rattachement des ventes et calcul des statuts clients'), force=True)
        self.env['yr.client.code'].sudo()._link_partners(cards)
        partner_ids = {partner.id for partner in self._existing_customers(cards).values()}
        partner_ids |= {partner.id for partner in self._customers_by_code(codes_in_file).values()}
        partner_ids |= self.env['yr.legacy.purchase']._link_partners(cards)
        SaleLine = self.env['yr.legacy.sale.line'].sudo()
        SaleLine._link_partners(cards)
        SaleLine._link_partners_by_code(codes_in_file)          # the client code is the finer link (shared cards)
        self._refresh_crm(partner_ids)

    def _customers_by_code(self, codes):
        found = {}
        Partner = self.env['res.partner'].with_context(active_test=False)
        for start in range(0, len(codes), 5000):
            for partner in Partner.search([('yr_client_code', 'in', codes[start:start + 5000])]):
                found.setdefault(partner.yr_client_code, partner)
        return found

    def _create_customers(self, Partner, batch, stats, errors):
        try:
            with self.env.cr.savepoint():
                Partner.create([vals for _line, vals in batch])
            stats['created'] += len(batch)
        except Exception:  # noqa: BLE001 - one bad line: retry one by one to report it
            for line, vals in batch:
                try:
                    with self.env.cr.savepoint():
                        Partner.create(vals)
                    stats['created'] += 1
                except Exception as error:  # noqa: BLE001
                    errors.append((line, vals.get('customer_code'), str(error).split('\n')[0]))

    def _import_client_new(self, stats, errors):
        self._import_customers('client_new', stats, errors)

    def _import_client_detail(self, stats, errors):
        self._import_customers('client_detail', stats, errors)

    def _import_client_old(self, stats, errors):
        self._import_customers('client_old', stats, errors)

    # ── 2-year turnover per customer: purchase history ───────────────────
    def _import_client_ca(self, stats, errors):
        """One fastmag purchase per card, day and store; re-importing updates the amounts."""
        visits = {}
        items = list(self._records('client_ca'))
        mapped = self.env['yr.client.code'].sudo()._cards_for(
            {client_code(i.get('fm_id')) for i in items if not card_number(i.get('card'))} - {''})
        for item in items:
            code = client_code(item.get('fm_id'))
            card, day = card_number(item.get('card')) or mapped.get(code, ''), parse_date(item.get('date'))
            if not card or not day:
                if card:
                    errors.append((item['_line'], card, _('date absente ou invalide')))
                elif day:
                    stats['anonymous'] += 1
                    self._note(item['_line'], code or text(item.get('card')), _('Non importé'),
                               _('pas de numéro de carte fidélité'))
                continue
            stats['rows'] += 1
            key = (card, datetime.combine(day, datetime.min.time()).replace(hour=12), text(item.get('shop')).upper())
            visit = visits.setdefault(key, {'amount': 0.0, 'client': code})
            visit['amount'] += number(item.get('amount')) or 0.0
        self._save_purchases(visits, stats, errors)

    # ── fastmag customer codes: Y… ↔ card ───────────────────────────────
    def _import_client_codes(self, stats, errors):
        """« Liste client » file: the code of the sales files (Y88074) and its loyalty card.

        Only a real card links the code to a customer; the others are kept (their sales stay anonymous).
        The sales imported before this file are attached to their customer now.
        """
        Code = self.env['yr.client.code'].sudo()
        items = {}
        for item in self._records('client_codes'):
            code = client_code(item.get('code'))
            if not code:
                if any(text(v) for k, v in item.items() if k != '_line'):
                    errors.append((item['_line'], text(item.get('code')), _('code client absent ou invalide')))
                continue
            stats['rows'] += 1
            card = card_number(item.get('card'))
            if not card:
                stats['codes_no_card'] += 1
                self._note(item['_line'], code, _('Gardé sans client'),
                           _('pas de vrai numéro de carte (%s) : ses ventes restent anonymes',
                             text(item.get('card')) or _('vide')))
            items[code] = {
                'code': code, 'card': card or False, 'name': text(item.get('name')) or False,
                'date': parse_date(item.get('date')), 'points': number(item.get('points')) or 0.0,
                'category': text(item.get('category')) or False, '_line': item['_line'],
            }
        stats['codes'] = len(items)
        existing = {}
        codes = list(items)
        for start in range(0, len(codes), 20000):
            for record in Code.search([('code', 'in', codes[start:start + 20000])]):
                existing[record.code] = record
        to_create, done = [], 0
        for code, vals in items.items():
            vals.pop('_line')
            record = existing.get(code)
            if record:
                changes = {k: v for k, v in vals.items() if (record[k] or False) != (v or False)}
                if changes:
                    record.write(changes)
                    stats['updated'] += 1
            else:
                to_create.append(vals)
            if len(to_create) >= 5000:
                stats['created'] += len(Code.create(to_create))
                to_create = []
                self._commit()
            done += 1
            self._progress(done, len(items), _('Enregistrement des codes clients'))
        if to_create:
            stats['created'] += len(Code.create(to_create))
        cards = [vals['card'] for vals in items.values() if vals['card']]
        Code._link_partners(cards)
        stats['codes_linked'] = Code.search_count([('code', 'in', codes), ('partner_id', '!=', False)]) \
            if len(codes) < 50000 else len(Code._codes_for(cards))
        # sales imported before: anonymous lines carrying a code with a card now
        SaleLine = self.env['yr.legacy.sale.line'].sudo()
        SaleLine.flush_model()
        self.env.cr.execute(SQL("""
            UPDATE yr_legacy_sale_line l SET card = c.card
              FROM yr_client_code c
             WHERE l.card IS NULL AND l.legacy_client_id = c.code AND c.card IS NOT NULL AND c.code = ANY(%s)
         RETURNING l.card
        """, codes))
        relinked = [card for (card,) in self.env.cr.fetchall()]
        SaleLine.invalidate_model(['card'])
        if relinked:
            stats['relinked'] = len(relinked)
            cards = list(set(relinked))
            SaleLine._link_partners(cards)
            self.env.cr.execute(SQL("""
                SELECT card, date, shop_code, sum(amount), min(legacy_client_id)
                  FROM yr_legacy_sale_line WHERE card = ANY(%s) GROUP BY card, date, shop_code
            """, cards))
            visits = {(card, datetime.combine(day, datetime.min.time()).replace(hour=12), shop):
                      {'amount': amount or 0.0, 'client': client}
                      for card, day, shop, amount, client in self.env.cr.fetchall()}
            self._save_purchases(visits, stats, errors, created='purchases', updated='purchases')

    def _save_purchases(self, visits, stats, errors, created='created', updated='updated', stores=None):
        """visits: {(card, datetime at noon, shop): {'amount', 'client'}} → yr.legacy.purchase (upsert)."""
        stores = stores if stores is not None else self._stores()
        unknown_shops = {shop for _card, _day, shop in visits if shop and shop not in stores}
        if unknown_shops:
            stats['stores_missing'] = unknown_shops
        Legacy = self.env['yr.legacy.purchase'].sudo()
        cards = list({card for card, _day, _shop in visits})
        existing = {}             # (card, day, shop) → (id, amount, store id): never whole records
        for chunk in range(0, len(cards), 20000):
            self.env.cr.execute(SQL("""
                SELECT card, date, COALESCE(shop_code, ''), id, amount, store_id
                  FROM yr_legacy_purchase WHERE card = ANY(%s)
            """, cards[chunk:chunk + 20000]))
            for card, day, shop, line_id, amount, store_id in self.env.cr.fetchall():
                existing[(card, day, shop)] = (line_id, amount or 0.0, store_id)
        to_create = []
        for (card, day, shop), visit in visits.items():
            store = stores.get(shop) if shop else False
            found = existing.get((card, day, shop or ''))
            if found:
                line_id, amount, store_id = found
                vals = {}
                if abs(amount - visit['amount']) > 0.0001:
                    vals['amount'] = visit['amount']
                if store and store_id != store.id:
                    vals['store_id'] = store.id
                if vals:
                    Legacy.browse(line_id).write(vals)
                    stats[updated] += 1
                continue
            if shop and not store:
                stats['unknown_stores'] += 1
            to_create.append({
                'card': card, 'date': day, 'shop_code': shop or False, 'store_id': store.id if store else False,
                'company_id': (store.company_id if store else self.env.company).id, 'amount': visit['amount'],
                'legacy_client_id': visit['client'] or False,
            })
            if len(to_create) >= 5000:
                stats[created] += len(Legacy.create(to_create))
                to_create = []
        if to_create:
            stats[created] += len(Legacy.create(to_create))
        Legacy._link_partners(cards)
        partner_ids = set()
        for chunk in range(0, len(cards), 5000):
            partner_ids |= set(Legacy.search([('card', 'in', cards[chunk:chunk + 5000]),
                                              ('partner_id', '!=', False)]).partner_id.ids)
        self._refresh_crm(partner_ids)
        stats['customers'] = len(set(Legacy.search([('card', 'in', cards), ('partner_id', '!=', False)])
                                     .mapped('partner_id').ids))
        unknown = len(set(Legacy.search([('card', 'in', cards), ('partner_id', '=', False)]).mapped('card')))
        if unknown:
            stats['unknown_cards'] = unknown
            errors.append(('-', '', _('%s carte(s) sans fiche client pour l\'instant : leurs achats sont gardés et '
                                      'seront rattachés à l\'import des fichiers clients.', unknown)))

    # ── store list: magasin, code ────────────────────────────────────────
    def _import_stores(self, stats, errors):
        """Creates the missing stores, sets the code of the others (matched on code, then on name)."""
        rows = self._rows(raw=True)
        first = next(rows)
        header = [norm(c) for c in first]
        fastmag_idx = None
        # optional « statut » / « fermé » / « actif » column: a closed store is kept, archived
        status_idx = next((header.index(h) for h in ('statut', 'etat', 'ferme', 'actif', 'status') if h in header),
                          None)
        if {'magasin', 'nummagasin'} <= set(header):
            # magasin | code fastmag (YR_CARR) | numéro (2): the number is the store code
            name_idx, code_idx = header.index('magasin'), header.index('nummagasin')
            fastmag_idx = next((header.index(h) for h in ('codemagasin', 'codemag') if h in header), None)
        elif {'magasin', 'code'} <= set(header):
            name_idx, code_idx = header.index('magasin'), header.index('code')
        else:       # no header row: name, code (, « fermé »)
            name_idx, code_idx, status_idx = 0, 1, 2
            rows = iter([first] + list(rows))
        status_col = header[status_idx] if status_idx is not None and status_idx < len(header) else ''

        def closed(value):
            value = norm(value)
            if status_col == 'actif':
                return value in ('non', 'n', '0', 'false', 'faux')
            return value.startswith(('ferm', 'clos', 'inactif', 'archiv')) or value in ('oui', 'o', '1', 'x') and \
                status_col == 'ferme'
        # every company: the stores already moved to their subsidiary are found whatever company is selected
        Store = self.env['loyalty.store'].sudo().with_context(active_test=False)
        stores = Store.search([])
        in_file = Store.browse()

        def code_taken(code, company, store=None):
            return any(s != store and s.company_id == company and (s.code or '').strip().upper() == code.upper()
                       for s in stores)
        for line_no, row in enumerate(rows, start=2):
            row = list(row) + [None, None, None]
            name, code = text(row[name_idx]), text(row[code_idx])
            fastmag = text(row[fastmag_idx]).upper() if fastmag_idx is not None else ''
            is_closed = status_idx is not None and closed(row[status_idx])
            if not name or not code:
                if name or code:
                    errors.append((line_no, name or code, _('nom du magasin ou code manquant')))
                continue
            stats['rows'] += 1
            store = stores.filtered(lambda s: (s.code or '').strip().upper() == code.upper()) or \
                stores.filtered(lambda s: norm(s.name) == norm(name))
            if len(store) > 1:
                errors.append((line_no, name, _('plusieurs magasins correspondent : %s',
                                                ', '.join(store.mapped('name')))))
                continue
            if store:
                in_file |= store
                vals = {}
                aliases = [a.strip() for a in (store.yr_fastmag_code or '').split(',') if a.strip()]
                if store.code != code:
                    if code_taken(code, store.company_id, store):
                        errors.append((line_no, name, _('le code %s est déjà celui d\'un autre magasin de la '
                                                        'société %s', code, store.company_id.name)))
                        continue
                    # the former code stays known for the files that still use it
                    vals['code'] = code
                    if store.code and store.code.upper() not in [x.upper() for x in aliases]:
                        aliases.append(store.code)
                if fastmag and fastmag not in [a.upper() for a in aliases]:
                    aliases.append(fastmag)                     # « YR_CARR » of the sales files
                if ','.join(aliases) != (store.yr_fastmag_code or ''):
                    vals['yr_fastmag_code'] = ','.join(aliases)
                if store.active == is_closed:
                    vals['active'] = not is_closed
                if vals:
                    store.write(vals)
                    stats['updated'] += 1
                else:
                    stats['skipped'] += 1
                continue
            web = any(word in norm(name) for word in ('eshop', 'siteweb', 'website', 'ecommerce'))
            if code_taken(code, self.env.company):
                errors.append((line_no, name, _('le code %s est déjà celui d\'un autre magasin de la société %s',
                                                code, self.env.company.name)))
                continue
            store = Store.create({'name': name, 'code': code, 'store_type': 'website' if web else 'store',
                                  'yr_fastmag_code': fastmag or False, 'active': not is_closed})
            stores |= store
            in_file |= store
            stats['created'] += 1
            if is_closed:
                stats['closed'] += 1
        if self.store_companies:
            self._store_companies_for(in_file, stats)
        if self.store_pos:
            self._store_pos_for(in_file, stats, errors)

    def _store_companies_for(self, stores, stats):
        """One company per store: a branch of the main company (found by name, else created). A store sharing its
        company with other stores, or in the main company itself, moves to its own branch."""
        Company = self.env['res.company'].sudo()
        Store = self.env['loyalty.store'].sudo().with_context(active_test=False)
        root = self.env.company.root_id or self.env.company
        family = Company.search([('id', 'child_of', root.id)])
        for rank, store in enumerate(stores.sudo(), start=1):
            self._progress(rank, len(stores), _('Sociétés (filiales) des magasins'))
            alone = Store.search_count([('company_id', '=', store.company_id.id), ('id', '!=', store.id)]) == 0
            if store.company_id != root and alone:
                continue                                    # already its own company
            company = family.filtered(lambda c: c != root and norm(c.name) == norm(store.name))[:1]
            if not company:
                name = store.name
                if Company.search_count([('name', '=', name)]):
                    name = '%s (%s)' % (store.name, store.code)  # company names are unique
                company = Company.create({'name': name, 'parent_id': root.id,
                                          'crm_is_website': store.store_type == 'website'})
                family |= company
                stats['companies'] += 1
            elif store.store_type == 'website' and not company.crm_is_website:
                company.crm_is_website = True
            if Store.search_count([('company_id', '=', company.id), ('code', '=', store.code), ('id', '!=', store.id)]):
                continue                                    # that subsidiary already has a store with this code
            store.company_id = company
            # the user who imports sees the new subsidiaries (others: Paramètres › Utilisateurs)
            if company not in self.env.user.company_ids:
                self.env.user.sudo().company_ids = [(4, company.id)]

    def _store_pos_for(self, stores, stats, errors):
        """A till (point de vente) for each open physical store without one, in the store's company: with a
        warehouse and the payment methods « Espèces » / « Carte bancaire » there if the company has none."""
        root = self.env.company.root_id or self.env.company
        todo = stores.sudo().filtered(lambda s: s.active and s.store_type == 'store' and not s.pos_config_ids)
        if not todo:
            return
        # new subsidiaries get their accounting (POS journal…) when the transaction is saved: now
        self._progress(0, len(todo), _('Comptabilité des filiales'), force=True)
        self.env.cr.precommit.run()
        for rank, store in enumerate(todo, start=1):
            self._progress(rank, len(todo), _('Caisses des magasins'), force=True)
            company = store.company_id
            env = self.env(context=dict(self.env.context, allowed_company_ids=[company.id] + root.ids),
                           su=True)
            try:
                with self.env.cr.savepoint():
                    Warehouse = env['stock.warehouse'].with_company(company)
                    if not Warehouse.search([('company_id', '=', company.id)], limit=1):
                        code = re.sub(r'[^A-Z0-9]', '', (store.name or '').upper())[:5] or 'MAG'
                        Warehouse.create({'name': store.name, 'code': code, 'company_id': company.id})
                    Method = env['pos.payment.method'].with_company(company)
                    methods = Method.search([('company_id', '=', company.id)])
                    if not methods:
                        Journal = env['account.journal']
                        domain = Journal._check_company_domain(company)
                        # a cash journal can serve one cash method only: the subsidiary gets its own
                        cash = Journal.search([('company_id', '=', company.id), ('type', '=', 'cash')], limit=1)
                        if not cash:
                            prefix = 'C' + re.sub(r'[^A-Z0-9]', '', (store.name or '').upper())[:3]
                            code = next(c for c in [prefix] + ['%s%d' % (prefix[:3], n) for n in range(1, 100)]
                                        if not Journal.search_count([*domain, ('code', '=', c)]))
                            cash = Journal.create({'name': _('Espèces %s', store.name), 'type': 'cash',
                                                   'code': code, 'company_id': company.id})
                        bank = Journal.search([*domain, ('type', '=', 'bank')], limit=1)
                        methods = Method.create([
                            {'name': _('Espèces'), 'journal_id': cash.id, 'company_id': company.id},
                            {'name': _('Carte bancaire'), 'journal_id': bank.id, 'company_id': company.id},
                        ])
                    env['pos.config'].with_company(company).create({
                        'name': store.name, 'crm_store_id': store.id, 'payment_method_ids': [(6, 0, methods.ids)]})
                    stats['pos'] += 1
            except Exception as error:  # noqa: BLE001 - reported, the other stores go on
                errors.append(('-', store.name, _('caisse non créée : %s', error)))

    # ── sellers: code magasin (first row of each store), vendeur, code logistique ──
    def _import_sellers(self, stats, errors):
        """Each seller becomes an employee of the company of its store, proposed at the store's tills."""
        stores = self._stores()
        index = self.env['yr.legacy.sale.line']._employee_index()
        Employee = self.env['hr.employee'].sudo().with_context(active_test=False)
        shop = ''
        for item in self._records('sellers'):
            shop = text(item.get('shop')).upper() or shop          # store code only on the first row
            name = text(item.get('seller'))
            if not name:
                continue
            stats['rows'] += 1
            store = stores.get(shop) if shop else None
            if shop and not store:
                # a store of the sellers file unknown in Posify: created (to rename / complete if needed)
                store_name = re.sub(r'^YR[_\s-]*', '', shop, flags=re.I) or shop
                store = self.env['loyalty.store'].sudo().create({
                    'name': store_name, 'code': shop, 'yr_fastmag_code': shop})
                stores.records |= store
                dict.__setitem__(stores, shop, store)
                stats['stores_created'] = stats.get('stores_created') or set()
                stats['stores_created'].add(store_name)
            code = text(item.get('logistic_code'))
            employee = Employee.search([('yr_logistic_code', '=', code)], limit=1) if code else Employee
            if not employee:
                same_name = index.get(' '.join(name.upper().split()))
                # same first name, other seller (other logistic code): a different person
                if same_name and (not code or not same_name.yr_logistic_code):
                    employee = same_name
            if employee:
                vals = {}
                if code and employee.yr_logistic_code != code:
                    vals['yr_logistic_code'] = code
                if store and employee.yr_store_id != store:
                    vals['yr_store_id'] = store.id
                if store and employee.company_id != store.company_id and not employee.user_id:
                    vals['company_id'] = store.company_id.id        # the company of its store
                if vals:
                    employee.write(vals)
                    stats['sellers_updated'] += 1
                else:
                    stats['skipped'] += 1
            else:
                employee = Employee.create({
                    'name': name, 'yr_logistic_code': code or False, 'yr_store_id': store.id if store else False,
                    'company_id': (store.company_id if store else self.env.company).id})
                index[' '.join(name.upper().split())] = employee
                stats['employees'] += 1
            if store:
                # proposed in the store's tills (« Vendeur avant paiement »)
                store.sudo().pos_config_ids.filtered(lambda c: employee not in c.yr_seller_employee_ids).write(
                    {'yr_seller_employee_ids': [(4, employee.id)]})
        self.env['yr.legacy.sale.line'].sudo()._link_employees()

    # ── sales detail: one line per ticket line, cashier and seller ───────
    def _deduced_reason(self, item):
        """No motif in the file: a free gift is told by its article code (Y…), a promotion by the discount."""
        if not self.deduce_motif or item.get('stamp'):
            return ''
        code = (text(item.get('ref')) or text(item.get('ean'))).upper()
        if code.startswith('Y'):
            return FREE_REASON
        if (number(item.get('discount')) or 0) >= 1:
            return DEDUCED_REASON
        return ''

    @staticmethod
    def _line_name(item, product):
        """Label of a sale line; an unknown article gets « name / barcode » (or reference) to be diagnosed."""
        name = text(item.get('name'))
        if product or item.get('stamp'):
            return name
        code = text(item.get('ean')) or text(item.get('ref'))
        return ' / '.join(part for part in (name, code) if part)

    def _import_sales_detail(self, stats, errors):
        SaleLine = self.env['yr.legacy.sale.line'].sudo()
        stores = self._stores()
        employees = SaleLine._employee_index()
        Employee = self.env['hr.employee'].sudo()
        Product = self.env['product.product'].with_context(active_test=False)
        items, refs, eans, tickets, rank = [], set(), set(), set(), defaultdict(int)
        for item in self._records('sales_detail'):
            shop, ticket, day = text(item.get('shop')).upper(), text(item.get('ticket')), parse_date(item.get('date'))
            if not shop and not ticket:
                continue
            if not shop or not ticket or not day:
                errors.append((item['_line'], ticket, _('magasin, n° de vente ou date manquant')))
                continue
            stats['rows'] += 1
            rank[(shop, ticket)] += 1               # position in the file: stable whatever lines are left out
            if is_coupon(item):
                stats['coupons'] += 1
                self._note(item['_line'], ticket, _('Non importé'), _(
                    'ligne de coupon / bon d\'achat (%s) : les coupons ne sont pas repris', text(item.get('ref'))))
                continue
            item.update(shop=shop, ticket=ticket, day=day, sequence=rank[(shop, ticket)])
            ref = text(item.get('ref'))
            item['stamp'] = ref.upper() in STAMP_REFS
            if not item['stamp']:
                refs.add(ref)
                eans.add(text(item.get('ean')))
            tickets.add((shop, ticket))
            items.append(item)
        stats['tickets'] = len(tickets)
        products = {}
        refs.discard('')
        eans.discard('')
        for product in Product.search(['|', ('default_code', 'in', list(refs)), ('barcode', 'in', list(eans))]):
            products.setdefault(('ref', product.default_code), product)
            products.setdefault(('ean', product.barcode), product)
        # lines already imported: only their key, id and order (never whole records: millions of lines)
        existing = defaultdict(list)                    # (shop, ticket) → [(sequence, id, order, article key)]
        keys = list(tickets)
        for start in range(0, len(keys), 20000):
            chunk = keys[start:start + 20000]
            self.env.cr.execute(SQL("""
                SELECT l.shop_code, l.ticket, l.sequence, l.id, l.pos_order_id, l.product_ref, l.barcode
                  FROM yr_legacy_sale_line l
                  JOIN unnest(%s::varchar[], %s::varchar[]) AS k(shop, ticket)
                    ON l.shop_code = k.shop AND l.ticket = k.ticket
                 ORDER BY l.sequence, l.id
            """, [shop for shop, _ticket in chunk], [ticket for _shop, ticket in chunk]))
            for shop, ticket, sequence, line_id, order_id, line_ref, line_ean in self.env.cr.fetchall():
                existing[(shop, ticket)].append((sequence, line_id, order_id, (line_ref or '', line_ean or '')))
        # a line is found again by its article (the nth time it comes in the ticket), not by its position: the
        # position moves when lines were left out (coupons) or missing at the first import
        free = {}                                       # (shop, ticket) → sequences taken
        pending = {}                                    # (shop, ticket, article key, nth) → existing line
        for ticket_key, lines in existing.items():
            free[ticket_key] = {line[0] for line in lines}
            nth = defaultdict(int)
            for line in lines:
                nth[line[3]] += 1
                pending[ticket_key + (line[3], nth[line[3]])] = line
                if line[3] == ('', ''):                 # nothing to recognise it by: found again by its position
                    pending[ticket_key + (line[0],)] = line
        ticket_has_order = {key for key, lines in existing.items() if any(line[2] for line in lines)}
        seen_nth = defaultdict(int)

        def employee(name, create, store):
            key = ' '.join(name.upper().split())
            if not key:
                return False
            if key not in employees and create:
                employees[key] = Employee.create({
                    'name': name, 'company_id': (store.company_id if store else self.env.company).id})
                stats['employees'] += 1
            return employees.get(key, Employee).id or False

        mapped = self.env['yr.client.code'].sudo()._cards_for(
            {client_code(i.get('fm_id')) for i in items if not card_number(i.get('card'))} - {''})
        unmapped, sellers_seen, cashiers_seen, unknown_refs = defaultdict(int), set(), set(), {}
        visits, to_create, cards, changed_tickets = {}, [], set(), set()
        codes_seen = set()
        done = 0
        for item in items:
            store = stores.get(item['shop'])
            if not store:
                stats['stores_missing'] = stats.get('stores_missing') or set()
                stats['stores_missing'].add(item['shop'])
            # the card of the file, else the card of the fastmag customer code (Y88074 / 88074)
            code = client_code(item.get('fm_id'))
            card = card_number(item.get('card'))
            if code:
                codes_seen.add(code)
            if not card and code:
                card = mapped.get(code, '')
                if card:
                    stats['via_code'] += 1
                else:
                    stats['unmapped_code'] += 1
                    unmapped[code] += 1
            elif not card:
                stats['anonymous'] += 1
            ref, ean = text(item.get('ref')), text(item.get('ean'))
            product = False if item['stamp'] else products.get(('ref', ref)) or products.get(('ean', ean))
            if not item['stamp'] and not product:
                stats['unknown_products'] += 1
                count = unknown_refs.get(ref or ean, (0,))[0]
                unknown_refs[ref or ean] = (count + 1, text(item.get('name')), ean)
            total = number(item.get('total')) or 0.0
            cashier, seller = text(item.get('cashier')), text(item.get('seller'))
            sellers_seen.add(seller.upper())
            cashiers_seen.add(cashier.upper())
            vals = {
                'shop_code': item['shop'], 'store_id': store.id if store else False,
                'pos_config_id': store.pos_config_ids[:1].id if store else False,
                'company_id': (store.company_id if store else self.env.company).id,
                'ticket': item['ticket'], 'sequence': item['sequence'], 'date': item['day'],
                'cashier_name': cashier or False, 'cashier_id': employee(cashier, False, store),
                'seller_name': seller or False, 'seller_id': employee(seller, self.create_sellers, store),
                'card': card or False, 'legacy_client_id': code or False,
                'product_ref': ref or False, 'barcode': ean or False, 'product_id': product.id if product else False,
                'name': self._line_name(item, product) or False, 'axe': text(item.get('axe')) or False,
                'product_line': text(item.get('line')) or False, 'is_stamp': item['stamp'],
                'qty': number(item.get('qty')) or 0.0, 'price_unit': number(item.get('price')) or 0.0,
                'discount': number(item.get('discount')) or 0.0, 'amount': total,
                'reason': text(item.get('reason')) or self._deduced_reason(item) or False, 'comment': text(item.get('comment')) or False,
            }
            ticket_key = (item['shop'], item['ticket'])
            article = (ref, ean)
            seen_nth[ticket_key + article] += 1
            found = pending.pop(ticket_key + (article, seen_nth[ticket_key + article]), None)
            if not found:
                found = pending.pop(ticket_key + (item['sequence'],), None)
            if found:
                vals['sequence'] = found[0]             # the line keeps its place
            else:
                taken = free.setdefault(ticket_key, set())
                if vals['sequence'] in taken:
                    vals['sequence'] = max(taken) + 1
                taken.add(vals['sequence'])
            if found and self.env.context.get('yr_resume'):
                stats['already'] += 1                   # saved before the interruption
            elif found:
                SaleLine.browse(found[1]).write(vals)
                stats['updated'] += 1
                if found[2]:
                    changed_tickets.add('%s/%s' % ticket_key)
                if stats['updated'] % 5000 == 0:
                    SaleLine.flush_model()
                    SaleLine.invalidate_model()         # keeps the memory flat on a big re-import
            else:
                to_create.append(vals)
                stats['added_to_known'] += 1 if ticket_key in existing else 0
                if ticket_key in ticket_has_order:      # a missing line: the POS order is made again with it
                    changed_tickets.add('%s/%s' % ticket_key)
            if card:
                cards.add(card)
                key = (card, datetime.combine(item['day'], datetime.min.time()).replace(hour=12), item['shop'])
                visit = visits.setdefault(key, {'amount': 0.0, 'client': code})
                visit['amount'] += total
            if len(to_create) >= 5000:
                stats['created'] += len(SaleLine.create(to_create))
                to_create = []
                self._commit()                          # an interruption resumes from here
            done += 1
            self._progress(done, len(items), _('Enregistrement des ventes'))
        if to_create:
            stats['created'] += len(SaleLine.create(to_create))
        stats['sellers_seen'] = len(sellers_seen - {''})
        stats['cashiers_seen'] = len(cashiers_seen - {''})
        for code, count in sorted(unmapped.items()):
            self._note('-', code, _('Vente sans client'), _(
                '%s ligne(s) : code client sans carte connue (importez la « Liste client » : elles seront '
                'rattachées)', count))
        for ref, (count, name, ean) in sorted(unknown_refs.items()):
            self._note('-', ref, _('Article inconnu'), _(
                '%(count)s ligne(s) gardée(s) sans article : %(name)s / %(code)s',
                count=count, name=name or '?', code=ean or ref))
        for shop in sorted(stats.get('stores_missing') or ()):
            self._note('-', shop, _('Magasin inconnu'), _('ventes gardées sans magasin'))
        if changed_tickets:
            # re-imported tickets: their POS orders are made again from the updated lines
            ticket_list = list(changed_tickets)
            for start in range(0, len(ticket_list), 20000):
                SaleLine._yr_delete_pos_orders(ticket_list[start:start + 20000])
        SaleLine._link_partners(list(cards))
        SaleLine._link_partners_by_code(codes_seen)             # customers without a card, or sharing one
        SaleLine.flush_model()
        SaleLine._link_motifs()
        if not self.env.context.get('yr_queue_item'):          # a queue does it once, after its last file
            self.env['loyalty.motif'].sudo()._infer_campaigns()
        if self.create_pos_orders:
            stats['pos_pending'] = len(tickets)
            SaleLine.flush_model()
            SaleLine._yr_trigger_pos_orders()
        if visits and self.update_purchases:
            self._save_purchases(visits, stats, errors, created='purchases', updated='purchases', stores=stores)

    def _refresh_crm(self, partner_ids):
        """Customer status, first / last purchase, turnover of Fidélité & CRM, now (normally hourly)."""
        if partner_ids:
            self.env['yr.legacy.sale.line']._yr_refresh_partners(partner_ids)
